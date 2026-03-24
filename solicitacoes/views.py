from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.core.paginator import Paginator
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
from django.db.models import Exists, OuterRef, Q
from .models import (
    Solicitacoes,
    SolicitacaoRotaItem,
    Recebedor,
    ClienteEmpresa,
    SolicitacaoTecnico,
    AuditoriaLog,
    EstornoHistorico,
    SolicitacaoExcluida,
)
from servicos.models import Servico
from django.utils import timezone
from datetime import time, datetime, timedelta
from io import BytesIO
import csv
import json
from usuarios.decorators import group_required, user_can_view_dashboard
from .audit_utils import serialize_instance

STATUS_VALIDOS = {choice[0] for choice in Solicitacoes._meta.get_field('status').choices}


def format_brl(value):
    """Formata número como moeda BRL: R$ 1.234,56 (2 decimais, ponto milhares, vírgula decimal)."""
    if value is None:
        return 'R$ 0,00'
    try:
        num = float(value)
    except (TypeError, ValueError):
        return 'R$ 0,00'
    s = f'{num:,.2f}'
    parts = s.split('.')
    int_part = parts[0].replace(',', '.')
    dec_part = parts[1] if len(parts) > 1 else '00'
    return f'R$ {int_part},{dec_part}'


def normalizar_status(valor):
    """
    Normaliza o valor do status para garantir consistência com os choices do modelo.
    Retorna 'pendente' como fallback seguro quando o valor for inválido.
    """
    if not valor:
        return 'pendente'
    valor_normalizado = valor.strip().lower()
    if valor_normalizado not in STATUS_VALIDOS:
        return 'pendente'
    return valor_normalizado


def registrar_exclusao_solicitacao(solicitacao, usuario=None, origem=None, motivo=None):
    """Cria registro na tabela de exclusões sem deletar o objeto principal."""
    SolicitacaoExcluida.objects.create(
        solicitacao=solicitacao,
        ticket=solicitacao.ticket,
        status=solicitacao.status,
        tipo=solicitacao.tipo,
        usuario=usuario if usuario and usuario.is_authenticated else None,
        origem=origem,
        motivo=motivo,
        dados_snapshot=serialize_instance(solicitacao),
    )


def garantir_log_movimentacao_card(request, solicitacao, status_anterior, status_novo, justificativa_estorno=''):
    """
    Garante que a movimentação de card seja registrada no AuditoriaLog.
    Se o signal já tiver registrado corretamente, não duplica.
    """
    if status_anterior == status_novo:
        return

    janela_inicio = timezone.now() - timedelta(seconds=10)
    log_recente = AuditoriaLog.objects.filter(
        app='solicitacoes',
        modelo='solicitacoes',
        objeto_id=str(solicitacao.pk),
        acao='update',
        origem=request.path,
        data_hora__gte=janela_inicio,
    ).order_by('-data_hora').first()

    if log_recente:
        antes = log_recente.dados_anteriores or {}
        depois = log_recente.dados_novos or {}
        if antes.get('status') == status_anterior and depois.get('status') == status_novo:
            return

    ip_address = request.META.get('HTTP_X_FORWARDED_FOR', '').split(',')[0].strip() or request.META.get('REMOTE_ADDR')
    user_agent = request.META.get('HTTP_USER_AGENT', '')

    dados_anteriores = {'status': status_anterior}
    dados_novos = {'status': status_novo}
    campos_alterados = ['status']

    if justificativa_estorno and status_novo == 'estorno':
        dados_novos['justificativa_estorno'] = justificativa_estorno
        campos_alterados.append('justificativa_estorno')

    AuditoriaLog.objects.create(
        app='solicitacoes',
        modelo='solicitacoes',
        objeto_id=str(solicitacao.pk),
        acao='update',
        usuario=request.user if request.user.is_authenticated else None,
        dados_anteriores=dados_anteriores,
        dados_novos=dados_novos,
        campos_alterados=campos_alterados,
        ip_address=ip_address,
        user_agent=user_agent,
        origem=request.path,
    )


def ping(request):
    """View simples para manter a sessão ativa."""
    return JsonResponse({'status': 'ok'})


def registrar_recebedor(nome, chave_pix=None):
    """
    Garante que o recebedor esteja registrado na tabela Recebedor.
    Se o recebedor já existir, NÃO altera chave_pix (evita sobrescrever com valor antigo da solicitação).
    Só define chave_pix ao CRIAR um novo recebedor.
    Retorna o objeto criado ou existente, ou None quando não foi possível registrar.
    """
    if not nome:
        return None
    
    chave_pix_normalizada = (chave_pix or "").strip() if chave_pix else ""
    nome_stripped = nome.strip()
    
    try:
        recebedor_obj = Recebedor.objects.filter(nome__iexact=nome_stripped).first()
        if recebedor_obj:
            # Não atualizar chave_pix do cadastro ao salvar solicitação (evita overwrite com valor antigo)
            return recebedor_obj
        # Só criar novo recebedor com chave_pix quando não existir
        recebedor_obj = Recebedor.objects.create(
            nome=nome_stripped,
            chave_pix=chave_pix_normalizada or "",
            ativo=True,
        )
        print(f"✅ Recebedor '{recebedor_obj.nome}' registrado automaticamente.")
        return recebedor_obj
    except Exception as e:
        print(f"⚠️ Erro ao registrar recebedor '{nome}': {e}")
        return None


def limpar_valor_monetario(valor_raw):
    """Função auxiliar para limpar e converter valores monetários formatados"""
    if not valor_raw:
        return 0.0
    try:
        # Remover formatação brasileira: R$, espaços, pontos de milhar, trocar vírgula por ponto
        valor_limpo = valor_raw.replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.').strip()
        return float(valor_limpo) if valor_limpo else 0.0
    except (ValueError, AttributeError):
        return 0.0

def garantir_migracao_campos():
    """
    Função auxiliar para garantir que os campos necessários existam no banco.
    Aplica a migração automaticamente se os campos não existirem.
    Também cria a tabela ClienteEmpresa se não existir.
    """
    try:
        from django.db import connection
        cursor = connection.cursor()
        
        campos_adicionados = False
        alteracoes_0007 = False

        # Verificar e criar tabela ClienteEmpresa se não existir
        cursor.execute("""
            SELECT name FROM sqlite_master 
            WHERE type='table' AND name='solicitacoes_clienteempresa';
        """)
        tabela_existe = cursor.fetchone()
        if not tabela_existe:
            try:
                cursor.execute("""
                    CREATE TABLE solicitacoes_clienteempresa (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        nome VARCHAR(200) NOT NULL UNIQUE,
                        cnpj VARCHAR(18),
                        ativo BOOLEAN NOT NULL DEFAULT 1,
                        data_criacao DATETIME NOT NULL,
                        data_atualizacao DATETIME NOT NULL
                    );
                """)
                print("✅ Tabela solicitacoes_clienteempresa criada automaticamente!")
                campos_adicionados = True
            except Exception as e:
                if "already exists" not in str(e).lower() and "duplicate" not in str(e).lower():
                    print(f"⚠️ Erro ao criar tabela ClienteEmpresa: {e}")
        else:
            # Se a tabela existe, verificar se tem o campo CNPJ
            cursor.execute("PRAGMA table_info(solicitacoes_clienteempresa)")
            columns_cliente = [row[1] for row in cursor.fetchall()]
            if 'cnpj' not in columns_cliente:
                try:
                    cursor.execute("ALTER TABLE solicitacoes_clienteempresa ADD COLUMN cnpj VARCHAR(18) DEFAULT NULL;")
                    print("✅ Campo cnpj adicionado à tabela solicitacoes_clienteempresa")
                    campos_adicionados = True
                except Exception as e:
                    if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                        print(f"⚠️ Erro ao adicionar campo cnpj na tabela ClienteEmpresa: {e}")

        # Verificar e criar tabela Recebedor se não existir
        cursor.execute("""
            SELECT name FROM sqlite_master 
            WHERE type='table' AND name='solicitacoes_recebedor';
        """)
        tabela_recebedor = cursor.fetchone()
        if not tabela_recebedor:
            try:
                cursor.execute("""
                    CREATE TABLE solicitacoes_recebedor (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        nome VARCHAR(100) NOT NULL UNIQUE,
                        chave_pix VARCHAR(255) NOT NULL,
                        ativo BOOLEAN NOT NULL DEFAULT 1,
                        data_criacao DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                        data_atualizacao DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
                    );
                """)
                print("✅ Tabela solicitacoes_recebedor criada automaticamente!")
                alteracoes_0007 = True
            except Exception as e:
                if "already exists" not in str(e).lower() and "duplicate" not in str(e).lower():
                    print(f"⚠️ Erro ao criar tabela Recebedor: {e}")
        else:
            # Se a tabela Recebedor existe, verificar se tem o campo supervisor
            cursor.execute("PRAGMA table_info(solicitacoes_recebedor)")
            columns_recebedor = [row[1] for row in cursor.fetchall()]
            if 'supervisor' not in columns_recebedor:
                try:
                    cursor.execute("""
                        ALTER TABLE solicitacoes_recebedor 
                        ADD COLUMN supervisor VARCHAR(50) DEFAULT NULL;
                    """)
                    print("✅ Campo supervisor adicionado à tabela solicitacoes_recebedor")
                    campos_adicionados = True
                except Exception as e:
                    if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                        print(f"⚠️ Erro ao adicionar campo supervisor na tabela Recebedor: {e}")
        
        cursor.execute("PRAGMA table_info(solicitacoes_solicitacoes)")
        columns = [row[1] for row in cursor.fetchall()]
        
        # Verificar se campos faltam e aplicar migração SQL diretamente
        campos_necessarios = {
            'data_entrada_status': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN data_entrada_status DATETIME DEFAULT NULL;",
            'valor_em_rota': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN valor_em_rota REAL DEFAULT 0.0;",
            'descricao_em_rota': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN descricao_em_rota TEXT DEFAULT NULL;",
            'chave_pix': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN chave_pix VARCHAR(255) DEFAULT NULL;",
            'cliente_empresa': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN cliente_empresa VARCHAR(200) DEFAULT NULL;",
            'cnpj': "ALTER TABLE solicitacoes_solicitacoes ADD COLUMN cnpj VARCHAR(18) DEFAULT NULL;"
        }
        
        campos_adicionados = False
        for campo, sql in campos_necessarios.items():
            if campo not in columns:
                try:
                    cursor.execute(sql)
                    print(f"✅ Campo {campo} adicionado ao banco de dados")
                    campos_adicionados = True
                except Exception as e:
                    if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                        print(f"⚠️ Erro ao adicionar campo {campo}: {e}")
        
        # Verificar também na tabela de itens de rota
        cursor.execute("PRAGMA table_info(solicitacoes_solicitacaorotaitem)")
        columns_rota = [row[1] for row in cursor.fetchall()]
        
        if 'cliente_empresa' not in columns_rota:
            try:
                cursor.execute("ALTER TABLE solicitacoes_solicitacaorotaitem ADD COLUMN cliente_empresa VARCHAR(200) DEFAULT NULL;")
                print(f"✅ Campo cliente_empresa adicionado à tabela solicitacoes_solicitacaorotaitem")
                campos_adicionados = True
            except Exception as e:
                if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                    print(f"⚠️ Erro ao adicionar campo cliente_empresa na tabela de itens: {e}")
        
        # Verificar se campo CNPJ existe na tabela de itens de rota
        if 'cnpj' not in columns_rota:
            try:
                cursor.execute("ALTER TABLE solicitacoes_solicitacaorotaitem ADD COLUMN cnpj VARCHAR(18) DEFAULT NULL;")
                print(f"✅ Campo cnpj adicionado à tabela solicitacoes_solicitacaorotaitem")
                campos_adicionados = True
            except Exception as e:
                if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                    print(f"⚠️ Erro ao adicionar campo cnpj na tabela de itens: {e}")

        # Verificar campos de recebedor na tabela de itens de rota (migração 0007)
        if 'recebedor' not in columns_rota:
            try:
                cursor.execute("ALTER TABLE solicitacoes_solicitacaorotaitem ADD COLUMN recebedor VARCHAR(100) DEFAULT NULL;")
                print("✅ Campo recebedor adicionado à tabela solicitacoes_solicitacaorotaitem")
                alteracoes_0007 = True
            except Exception as e:
                if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                    print(f"⚠️ Erro ao adicionar campo recebedor na tabela de itens: {e}")

        if 'chave_pix' not in columns_rota:
            try:
                cursor.execute("ALTER TABLE solicitacoes_solicitacaorotaitem ADD COLUMN chave_pix VARCHAR(255) DEFAULT NULL;")
                print("✅ Campo chave_pix adicionado à tabela solicitacoes_solicitacaorotaitem")
                alteracoes_0007 = True
            except Exception as e:
                if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                    print(f"⚠️ Erro ao adicionar campo chave_pix na tabela de itens: {e}")
        
        # Verificar se campo CNPJ existe na tabela principal
        if 'cnpj' not in columns:
            try:
                cursor.execute("ALTER TABLE solicitacoes_solicitacoes ADD COLUMN cnpj VARCHAR(18) DEFAULT NULL;")
                print(f"✅ Campo cnpj adicionado à tabela solicitacoes_solicitacoes")
                campos_adicionados = True
            except Exception as e:
                if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                    print(f"⚠️ Erro ao adicionar campo cnpj: {e}")
        
        # Verificar tabela SolicitacaoTecnico e seus campos
        cursor.execute("""
            SELECT name FROM sqlite_master 
            WHERE type='table' AND name='solicitacoes_solicitacaotecnico';
        """)
        tabela_tecnico = cursor.fetchone()
        if tabela_tecnico:
            cursor.execute("PRAGMA table_info(solicitacoes_solicitacaotecnico)")
            columns_tecnico = [row[1] for row in cursor.fetchall()]
            
            # Verificar e adicionar campo cliente_empresa_id se não existir
            if 'cliente_empresa_id' not in columns_tecnico:
                try:
                    cursor.execute("ALTER TABLE solicitacoes_solicitacaotecnico ADD COLUMN cliente_empresa_id INTEGER DEFAULT NULL REFERENCES solicitacoes_clienteempresa(id);")
                    print("✅ Campo cliente_empresa_id adicionado à tabela solicitacoes_solicitacaotecnico")
                    campos_adicionados = True
                except Exception as e:
                    if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                        print(f"⚠️ Erro ao adicionar campo cliente_empresa_id na tabela SolicitacaoTecnico: {e}")
            
            # Verificar e adicionar campo data_pagamento se não existir
            if 'data_pagamento' not in columns_tecnico:
                try:
                    cursor.execute("ALTER TABLE solicitacoes_solicitacaotecnico ADD COLUMN data_pagamento DATE DEFAULT NULL;")
                    print("✅ Campo data_pagamento adicionado à tabela solicitacoes_solicitacaotecnico")
                    campos_adicionados = True
                except Exception as e:
                    if "duplicate column" not in str(e).lower() and "already exists" not in str(e).lower():
                        print(f"⚠️ Erro ao adicionar campo data_pagamento na tabela SolicitacaoTecnico: {e}")
        
        # Se data_entrada_status foi adicionado, atualizar registros existentes
        if campos_adicionados:
            try:
                # Verificar se o campo existe agora (pode ter sido adicionado)
                cursor.execute("PRAGMA table_info(solicitacoes_solicitacoes)")
                columns_atualizadas = [row[1] for row in cursor.fetchall()]
                if 'data_entrada_status' in columns_atualizadas:
                    cursor.execute("""
                        UPDATE solicitacoes_solicitacoes 
                        SET data_entrada_status = datetime(data_de_criacao || ' ' || time(tempo_criacao))
                        WHERE data_entrada_status IS NULL;
                    """)
                    print(f"✅ {cursor.rowcount} registros atualizados com data_entrada_status")
            except Exception as e:
                print(f"⚠️ Erro ao atualizar data_entrada_status: {e}")
        
        # Registrar migração se aplicou alguma
        if campos_adicionados:
            try:
                cursor.execute(
                    "INSERT OR IGNORE INTO django_migrations (app, name, applied) VALUES (?, ?, ?)",
                    ['solicitacoes', '0006_add_data_entrada_status', timezone.now()]
                )
                print("✅ Migração 0006 registrada no Django")
            except Exception as e:
                print(f"⚠️ Erro ao registrar migração (pode ser ignorado): {e}")
        if alteracoes_0007:
            try:
                cursor.execute(
                    "INSERT OR IGNORE INTO django_migrations (app, name, applied) VALUES (?, ?, ?)",
                    ['solicitacoes', '0007_recebedor_solicitacaorotaitem_chave_pix_and_more', timezone.now()]
                )
                print("✅ Migração 0007 registrada no Django")
            except Exception as e:
                print(f"⚠️ Erro ao registrar migração 0007 (pode ser ignorado): {e}")
    except Exception as e:
        print(f"⚠️ Erro ao verificar/aplicar migração: {e}")

def receber_dados(request):
    # Garantir que o usuário está autenticado antes de qualquer processamento
    if not request.user.is_authenticated:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'message': 'Sessão expirada. Por favor, faça login novamente.'}, status=401)
        return redirect('/usuarios/login/')

    # Garantir que a migração está aplicada antes de processar qualquer requisição
    try:
        garantir_migracao_campos()
    except Exception as e:
        print(f"⚠️ Aviso ao garantir migração: {e}")
    garantir_migracao_campos()
    
    if request.method == 'POST':
        try:
            solicitacao_id_raw = (request.POST.get('solicitacao_id') or '').strip()
            print(f"🔍 DEBUG EDIÇÃO - solicitacao_id recebido: '{solicitacao_id_raw}'")
            print(f"🔍 DEBUG EDIÇÃO - Todos os campos POST recebidos: {list(request.POST.keys())}")
            solicitacao_existente = None
            if solicitacao_id_raw:
                try:
                    solicitacao_existente = Solicitacoes.objects.select_related('servico').prefetch_related('itens_rota').get(id=int(solicitacao_id_raw))
                    print(f"✅ DEBUG EDIÇÃO - Solicitação encontrada: ID={solicitacao_existente.id}, Ticket={solicitacao_existente.ticket}, Status={solicitacao_existente.status}")
                except (ValueError, Solicitacoes.DoesNotExist) as e:
                    print(f"❌ DEBUG EDIÇÃO - Erro ao buscar solicitação: {e}")
                    messages.error(request, 'Solicitação para edição não foi encontrada ou já foi removida.')
                    return redirect('/solicitacoes/home/')
                
                if normalizar_status(solicitacao_existente.status) != 'pendente':
                    print(f"❌ DEBUG EDIÇÃO - Status não permite edição: {solicitacao_existente.status}")
                    messages.error(request, 'Somente solicitações pendentes podem ser editadas.')
                    return redirect('/solicitacoes/home/')
            
            is_edit_mode = solicitacao_existente is not None
            print(f"🔍 DEBUG EDIÇÃO - Modo edição ativado: {is_edit_mode}")
            success_message = 'Solicitação atualizada com sucesso!' if is_edit_mode else 'Solicitação criada com sucesso!'
            
            # Verificar se é uma solicitação de técnico ANTES de processar outros tipos
            # IMPORTANTE: Verificar não apenas se o campo existe, mas se está PREENCHIDO
            tecnico_ids_encontrados = []
            modo_tecnico = request.POST.get('modo_tecnico', '').strip().lower()
            
            # Buscar campos de técnico no POST E verificar se estão PREENCHIDOS
            # IMPORTANTE: Verificar se TODOS os campos obrigatórios estão preenchidos, não apenas o ID
            for key in request.POST.keys():
                if key.startswith('tecnico_solicitacao_'):
                    try:
                        num_id = int(key.replace('tecnico_solicitacao_', ''))
                        # Verificar se o campo tem valor (não está vazio)
                        valor_campo = request.POST.get(key, '').strip()
                        if valor_campo:  # Se o ID existe, verificar se outros campos também estão preenchidos
                            # Verificar se há recebedor, serviço e valor para este ID
                            recebedor = request.POST.get(f'tecnico_recebedor_{num_id}', '').strip()
                            servico = request.POST.get(f'tecnico_servico_{num_id}', '').strip()
                            valor = request.POST.get(f'tecnico_valor_pagamento_{num_id}', '').strip()
                            data_realizacao = request.POST.get(f'tecnico_data_realizacao_{num_id}', '').strip()
                            
                            # Só adicionar se TODOS os campos obrigatórios estiverem preenchidos
                            if recebedor and servico and valor and data_realizacao:
                                tecnico_ids_encontrados.append(num_id)
                                print(f"✅ ID {num_id} válido - todos os campos obrigatórios preenchidos")
                            else:
                                print(f"⚠️ ID {num_id} ignorado - campos obrigatórios faltando (recebedor={bool(recebedor)}, servico={bool(servico)}, valor={bool(valor)}, data={bool(data_realizacao)})")
                    except ValueError:
                        continue
            
            # Verificar também se há outros campos de técnico preenchidos para garantir
            # Só considerar técnico se modo_tecnico for explicitamente 'true' E houver IDs preenchidos
            # OU se houver pelo menos um ID de técnico com todos os campos obrigatórios preenchidos
            is_solicitacao_tecnico = False
            
            # Se modo_tecnico é 'true', verificar se realmente há dados válidos
            if modo_tecnico == 'true' and len(tecnico_ids_encontrados) > 0:
                # Verificar se pelo menos um item de técnico tem todos os campos obrigatórios
                for i in tecnico_ids_encontrados:
                    tecnico_recebedor = request.POST.get(f'tecnico_recebedor_{i}', '').strip()
                    tecnico_servico_id = request.POST.get(f'tecnico_servico_{i}', '').strip()
                    tecnico_valor_pagamento = request.POST.get(f'tecnico_valor_pagamento_{i}', '').strip()
                    tecnico_data_realizacao = request.POST.get(f'tecnico_data_realizacao_{i}', '').strip()
                    
                    # Se pelo menos um item tem todos os campos obrigatórios, é técnico
                    if tecnico_recebedor and tecnico_servico_id and tecnico_valor_pagamento and tecnico_data_realizacao:
                        is_solicitacao_tecnico = True
                        break
            
            # Se modo_tecnico não é 'true', NÃO é solicitação de técnico (mesmo que haja campos vazios)
            if modo_tecnico != 'true':
                is_solicitacao_tecnico = False
                print(f"🔍 DEBUG: modo_tecnico não é 'true' ({modo_tecnico}), ignorando campos de técnico")
            
            print(f"🔍 DEBUG: Verificando solicitação de técnico - IDs encontrados: {tecnico_ids_encontrados}, modo_tecnico: {modo_tecnico}, is_solicitacao_tecnico: {is_solicitacao_tecnico}")
            print(f"🔍 DEBUG: Primeiros 50 campos POST: {list(request.POST.keys())[:50]}")
            
            if is_solicitacao_tecnico:
                # Processar solicitação de técnico (múltiplos tickets)
                print(f"🔍 DEBUG: Processando solicitação de técnico - IDs encontrados: {tecnico_ids_encontrados}")
                
                # Ordenar para processar na ordem correta
                tecnico_ids_encontrados.sort()
                
                status = normalizar_status(request.POST.get('status') or 'pendente')
                agora = timezone.now()
                data_aprovacao_inicial = None
                if status in ['aprovado', 'concluido']:
                    data_aprovacao_inicial = agora
                
                solicitacoes_criadas = []
                itens_tecnico = []  # Armazenar todos os itens antes de criar a solicitação principal
                valor_total_geral = 0.0
                tickets_tecnico = []  # Armazenar todos os tickets para o título
                primeiro_recebedor = ''
                primeira_chave_pix = ''
                
                try:
                    with transaction.atomic():
                        # Primeiro, validar e coletar todos os dados dos itens
                        for i in tecnico_ids_encontrados:
                            # Capturar dados do formulário de técnico para este ID
                            tecnico_id_solicitacao = request.POST.get(f'tecnico_solicitacao_{i}', '').strip()
                            tecnico_recebedor = request.POST.get(f'tecnico_recebedor_{i}', '').strip()
                            tecnico_pix = request.POST.get(f'tecnico_pix_{i}', '').strip()
                            tecnico_servico_id = request.POST.get(f'tecnico_servico_{i}', '').strip()
                            tecnico_cliente_empresa_id = request.POST.get(f'tecnico_cliente_empresa_{i}', '').strip()
                            tecnico_valor_pagamento_raw = request.POST.get(f'tecnico_valor_pagamento_{i}', '').strip()
                            tecnico_valor_extra_raw = request.POST.get(f'tecnico_valor_extra_{i}', '').strip()
                            tecnico_descricao = request.POST.get(f'tecnico_descricao_{i}', '').strip()
                            tecnico_data_realizacao_str = request.POST.get(f'tecnico_data_realizacao_{i}', '').strip()
                            tecnico_data_pagamento_str = request.POST.get(f'tecnico_data_pagamento_{i}', '').strip()
                            tecnico_atividade_produtiva = request.POST.get(f'tecnico_atividade_produtiva_{i}', 'true').strip().lower() == 'true'
                            
                            # Verificar se é requisição AJAX para retornar JSON em caso de erro
                            is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
                            
                            # Validações
                            if not tecnico_recebedor:
                                error_msg = f'O campo "Nome do Técnico" do ID {i} é obrigatório.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            if not tecnico_id_solicitacao:
                                error_msg = f'O campo "ID da Solicitação" do ID {i} é obrigatório.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            if not tecnico_servico_id:
                                error_msg = f'O campo "Tipo de Serviço" do ID {i} é obrigatório.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            if not tecnico_data_realizacao_str:
                                error_msg = f'O campo "Data da Realização da Atividade" do ID {i} é obrigatório.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            # Buscar recebedor
                            try:
                                recebedor_obj = Recebedor.objects.get(nome=tecnico_recebedor)
                            except Recebedor.DoesNotExist:
                                error_msg = f'Recebedor "{tecnico_recebedor}" do ID {i} não encontrado no sistema.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            # Buscar serviço
                            try:
                                servico_obj = Servico.objects.get(id=tecnico_servico_id, ativo=True)
                            except Servico.DoesNotExist:
                                error_msg = f'Serviço selecionado do ID {i} não foi encontrado ou está inativo.'
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            # Buscar cliente/empresa (opcional)
                            cliente_empresa_obj = None
                            if tecnico_cliente_empresa_id:
                                print(f"🔍 DEBUG Item {i}: Buscando Cliente/Empresa ID: '{tecnico_cliente_empresa_id}'")
                                try:
                                    cliente_empresa_obj = ClienteEmpresa.objects.get(id=int(tecnico_cliente_empresa_id), ativo=True)
                                    print(f"✅ DEBUG Item {i}: Cliente/Empresa encontrado: {cliente_empresa_obj.nome}")
                                except ClienteEmpresa.DoesNotExist:
                                    # Não é erro crítico, apenas log
                                    print(f"⚠️ Cliente/Empresa ID {tecnico_cliente_empresa_id} não encontrado ou inativo para o ID {i}")
                                except (ValueError, TypeError) as e:
                                    print(f"⚠️ ID de Cliente/Empresa inválido para o ID {i}: '{tecnico_cliente_empresa_id}', Erro: {e}")
                            
                            # Converter valores monetários
                            print(f"🔍 DEBUG Item {i}: Valor pagamento raw: '{tecnico_valor_pagamento_raw}', Valor extra raw: '{tecnico_valor_extra_raw}'")
                            valor_pagamento_tecnico = limpar_valor_monetario(tecnico_valor_pagamento_raw)
                            valor_extra = limpar_valor_monetario(tecnico_valor_extra_raw) if tecnico_valor_extra_raw else 0.0
                            print(f"🔍 DEBUG Item {i}: Valor pagamento convertido: {valor_pagamento_tecnico}, Valor extra convertido: {valor_extra}")
                            
                            if valor_pagamento_tecnico <= 0:
                                error_msg = f'O valor a pagar para o técnico do ID {i} deve ser maior que zero. Valor recebido: "{tecnico_valor_pagamento_raw}"'
                                print(f"❌ ERRO: {error_msg}")
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            # Validar e converter data de realização
                            try:
                                data_realizacao_atividade = datetime.strptime(tecnico_data_realizacao_str, '%Y-%m-%d').date()
                                print(f"🔍 DEBUG Item {i}: Data realização convertida: {data_realizacao_atividade}")
                            except (ValueError, TypeError) as e:
                                error_msg = f'Data da realização da atividade do ID {i} inválida. Use o formato correto. Valor recebido: "{tecnico_data_realizacao_str}"'
                                print(f"❌ ERRO: {error_msg}, Exception: {e}")
                                if is_ajax:
                                    return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                messages.error(request, error_msg)
                                return redirect('/solicitacoes/home/')
                            
                            # Validar e converter data de pagamento (opcional)
                            data_pagamento = None
                            if tecnico_data_pagamento_str:
                                try:
                                    data_pagamento = datetime.strptime(tecnico_data_pagamento_str, '%Y-%m-%d').date()
                                    print(f"🔍 DEBUG Item {i}: Data pagamento convertida: {data_pagamento}")
                                except (ValueError, TypeError) as e:
                                    error_msg = f'Data de pagamento do ID {i} inválida. Use o formato correto. Valor recebido: "{tecnico_data_pagamento_str}"'
                                    print(f"❌ ERRO: {error_msg}, Exception: {e}")
                                    if is_ajax:
                                        return JsonResponse({'success': False, 'message': error_msg}, status=400)
                                    messages.error(request, error_msg)
                                    return redirect('/solicitacoes/home/')
                            
                            valor_total = valor_pagamento_tecnico + valor_extra
                            valor_total_geral += valor_total
                            
                            # Armazenar primeiro recebedor e chave PIX para a solicitação principal
                            if not primeiro_recebedor:
                                primeiro_recebedor = tecnico_recebedor
                            if not primeira_chave_pix:
                                primeira_chave_pix = tecnico_pix or recebedor_obj.chave_pix
                            
                            # Armazenar ticket para o título
                            tickets_tecnico.append(tecnico_id_solicitacao)
                            
                            # Armazenar dados do item para criar depois
                            itens_tecnico.append({
                                'tecnico_id_solicitacao': tecnico_id_solicitacao,
                                'recebedor_obj': recebedor_obj,
                                'tecnico_pix': tecnico_pix or recebedor_obj.chave_pix,
                                'servico_obj': servico_obj,
                                'cliente_empresa_obj': cliente_empresa_obj,
                                'valor_pagamento_tecnico': valor_pagamento_tecnico,
                                'valor_extra': valor_extra,
                                'tecnico_descricao': tecnico_descricao,
                                'data_realizacao_atividade': data_realizacao_atividade,
                                'data_pagamento': data_pagamento,
                                'atividade_produtiva': tecnico_atividade_produtiva,
                                'valor_total': valor_total
                            })
                            
                            print(f"✅ Item de técnico {i} validado e preparado. Ticket: {tecnico_id_solicitacao}, Valor: {valor_total}")
                        
                        # Em solicitação agrupada por técnico, todos os itens devem ter o MESMO recebedor e chave PIX do primeiro item
                        if len(itens_tecnico) >= 1:
                            primeiro_recebedor_obj = itens_tecnico[0]['recebedor_obj']
                            primeira_chave_pix = (itens_tecnico[0].get('tecnico_pix') or '') or (getattr(primeiro_recebedor_obj, 'chave_pix', None) or '')
                            primeiro_recebedor = primeiro_recebedor_obj.nome
                            for item in itens_tecnico:
                                item['recebedor_obj'] = primeiro_recebedor_obj
                                item['tecnico_pix'] = primeira_chave_pix
                        
                        # Agora criar APENAS UMA solicitação principal
                        if len(itens_tecnico) == 0:
                            error_msg = 'Nenhum item de técnico válido encontrado.'
                            if is_ajax:
                                return JsonResponse({'success': False, 'message': error_msg}, status=400)
                            messages.error(request, error_msg)
                            return redirect('/solicitacoes/home/')
                        
                        print(f"🔍 DEBUG: Criando UMA ÚNICA solicitação principal com {len(itens_tecnico)} item(ns)")
                        print(f"🔍 DEBUG: Tickets coletados: {tickets_tecnico}")
                        
                        # Gerar ticket principal (usar o primeiro ticket ou criar um ticket único)
                        ticket_principal = tickets_tecnico[0]
                        
                        # Verificar se o ticket já existe
                        ticket_existe = Solicitacoes.objects.filter(ticket=ticket_principal).exists()
                        if ticket_existe:
                            # Se o ticket principal já existe, gerar um ticket único com timestamp
                            timestamp = int(timezone.now().timestamp())
                            ticket_principal = f"TECNICO-{timestamp}"
                            print(f"⚠️ Ticket '{tickets_tecnico[0]}' já existe. Usando ticket único: '{ticket_principal}'")
                        else:
                            print(f"✅ Ticket principal '{ticket_principal}' é único - usando este ticket")
                        
                        # Criar título com todos os tickets (similar a Em Rota)
                        # Limitar o título a 70 caracteres (limite do campo no modelo)
                        if len(tickets_tecnico) > 3:
                            titulo_principal = f"Solicitação de Técnico - {', '.join(tickets_tecnico[:3])}... (+{len(tickets_tecnico) - 3})"
                        else:
                            titulo_principal = f"Solicitação de Técnico - {', '.join(tickets_tecnico)}"
                        
                        # Garantir que o título não ultrapasse 70 caracteres
                        if len(titulo_principal) > 70:
                            titulo_principal = titulo_principal[:67] + "..."
                        
                        print(f"🔍 DEBUG: Ticket principal: '{ticket_principal}', Título: '{titulo_principal}', Valor total: R$ {valor_total_geral:.2f}")
                        
                        # Criar descrição combinada ou usar a primeira
                        descricao_principal = ''
                        if itens_tecnico[0]['tecnico_descricao']:
                            descricao_principal = itens_tecnico[0]['tecnico_descricao']
                        else:
                            descricao_principal = f'Solicitação de técnico com {len(itens_tecnico)} item(ns)'
                        
                        # Capturar anexos e prioridade para técnico
                        anexo_tecnico = request.FILES.get('tecnico_anexos')
                        prioridade_tecnico = request.POST.get('tecnico_priority', 'baixa')
                        
                        if is_edit_mode:
                            # MODO EDIÇÃO: Atualizar solicitação existente
                            solicitacao_principal = solicitacao_existente
                            print(f"🔍 DEBUG EDIÇÃO TÉCNICO - Atualizando solicitação ID={solicitacao_principal.id}")
                            
                            # Atualizar campos da solicitação principal
                            solicitacao_principal.titulo = titulo_principal
                            solicitacao_principal.nome_do_recebedor = primeiro_recebedor
                            solicitacao_principal.chave_pix = primeira_chave_pix
                            solicitacao_principal.valor = valor_total_geral
                            solicitacao_principal.descricao = descricao_principal
                            solicitacao_principal.prioridade = prioridade_tecnico
                            solicitacao_principal.servico = itens_tecnico[0]['servico_obj']
                            if anexo_tecnico:
                                solicitacao_principal.anexo = anexo_tecnico
                            solicitacao_principal.save()
                            
                            print(f"✅ Solicitação principal atualizada: ID={solicitacao_principal.id}, Ticket={solicitacao_principal.ticket}, Valor Total={valor_total_geral}")
                            
                            # Deletar todos os itens de técnico existentes
                            itens_antigos = SolicitacaoTecnico.objects.filter(solicitacao=solicitacao_principal)
                            num_itens_antigos = itens_antigos.count()
                            itens_antigos.delete()
                            print(f"🗑️ {num_itens_antigos} item(ns) de técnico antigo(s) deletado(s)")
                            
                            # Criar novos itens de técnico com os dados atualizados
                            for idx, item in enumerate(itens_tecnico):
                                solicitacao_tecnico = SolicitacaoTecnico.objects.create(
                                    solicitacao=solicitacao_principal,
                                    ticket_item=item['tecnico_id_solicitacao'],
                                    recebedor=item['recebedor_obj'],
                                    servico=item['servico_obj'],
                                    cliente_empresa=item.get('cliente_empresa_obj'),
                                    valor_pagamento_tecnico=item['valor_pagamento_tecnico'],
                                    valor_extra=item['valor_extra'],
                                    descricao=item['tecnico_descricao'],
                                    data_realizacao_atividade=item['data_realizacao_atividade'],
                                    data_pagamento=item.get('data_pagamento'),
                                    atividade_produtiva=item['atividade_produtiva']
                                )
                                
                                solicitacoes_criadas.append({
                                    'principal_id': solicitacao_principal.id,
                                    'tecnico_id': solicitacao_tecnico.id,
                                    'ticket': item['tecnico_id_solicitacao']
                                })
                                
                                print(f"✅ Item de técnico {idx + 1}/{len(itens_tecnico)} atualizado! ID Técnico: {solicitacao_tecnico.id}, Ticket: {item['tecnico_id_solicitacao']}")
                        else:
                            # MODO CRIAÇÃO: Criar nova solicitação
                            solicitacao_principal = Solicitacoes.objects.create(
                                ticket=ticket_principal,
                                status=status,
                                titulo=titulo_principal,
                                nome_solicitante=request.user,
                                nome_do_recebedor=primeiro_recebedor,
                                chave_pix=primeira_chave_pix,
                                valor=valor_total_geral,
                                descricao=descricao_principal,
                                data_de_pagamento=agora.date(),
                                data_de_criacao=agora.date(),
                                tempo_criacao=agora.time(),
                                tempo_fila=timezone.now().time().replace(hour=0, minute=0, second=0, microsecond=0),
                                data_entrada_status=agora,
                                data_aprovacao=data_aprovacao_inicial,
                                prioridade=prioridade_tecnico,
                                anexo=anexo_tecnico,
                                servico=itens_tecnico[0]['servico_obj'],
                                tipo='casual'  # Usar tipo casual para aparecer no Kanban
                            )
                            
                            print(f"✅ Solicitação principal criada: ID={solicitacao_principal.id}, Ticket={ticket_principal}, Valor Total={valor_total_geral}")
                            
                            # Criar todos os registros de SolicitacaoTecnico vinculados à solicitação principal
                            for idx, item in enumerate(itens_tecnico):
                                solicitacao_tecnico = SolicitacaoTecnico.objects.create(
                                    solicitacao=solicitacao_principal,
                                    ticket_item=item['tecnico_id_solicitacao'],
                                    recebedor=item['recebedor_obj'],
                                    servico=item['servico_obj'],
                                    cliente_empresa=item.get('cliente_empresa_obj'),
                                    valor_pagamento_tecnico=item['valor_pagamento_tecnico'],
                                    valor_extra=item['valor_extra'],
                                    descricao=item['tecnico_descricao'],
                                    data_realizacao_atividade=item['data_realizacao_atividade'],
                                    data_pagamento=item.get('data_pagamento'),
                                    atividade_produtiva=item['atividade_produtiva']
                                )
                                
                                solicitacoes_criadas.append({
                                    'principal_id': solicitacao_principal.id,
                                    'tecnico_id': solicitacao_tecnico.id,
                                    'ticket': item['tecnico_id_solicitacao']
                                })
                                
                                print(f"✅ Item de técnico {idx + 1}/{len(itens_tecnico)} criado! ID Técnico: {solicitacao_tecnico.id}, Ticket: {item['tecnico_id_solicitacao']}, Vinculado à solicitação principal ID={solicitacao_principal.id}")
                        
                        # Verificação final: garantir que todos os itens estão vinculados à mesma solicitação principal
                        itens_criados = SolicitacaoTecnico.objects.filter(solicitacao=solicitacao_principal)
                        acao = 'atualizado(s)' if is_edit_mode else 'criado(s)'
                        print(f"✅ Total: {len(solicitacoes_criadas)} item(ns) de técnico {acao} vinculado(s) à solicitação principal ID={solicitacao_principal.id}")
                        print(f"🔍 DEBUG: Verificação final - {itens_criados.count()} itens encontrados vinculados à solicitação principal ID={solicitacao_principal.id}")
                        
                        # Confirmar que apenas UMA solicitação foi criada (apenas em modo criação)
                        if not is_edit_mode:
                            solicitacoes_com_ticket = Solicitacoes.objects.filter(ticket=ticket_principal)
                            if solicitacoes_com_ticket.count() > 1:
                                print(f"❌ ERRO: Foram criadas {solicitacoes_com_ticket.count()} solicitações com o mesmo ticket '{ticket_principal}'! Isso não deveria acontecer.")
                            else:
                                print(f"✅ Confirmado: Apenas 1 solicitação foi criada com o ticket '{ticket_principal}'")
                        
                        # Verificar se é requisição AJAX
                        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
                        if is_ajax:
                            mensagem = f'1 solicitação de técnico atualizada com {len(solicitacoes_criadas)} item(ns) com sucesso!' if is_edit_mode else f'1 solicitação de técnico criada com {len(solicitacoes_criadas)} item(ns) com sucesso!'
                            return JsonResponse({
                                'success': True,
                                'message': mensagem
                            })
                        
                        mensagem = f'1 solicitação de técnico atualizada com {len(solicitacoes_criadas)} item(ns) com sucesso!' if is_edit_mode else f'1 solicitação de técnico criada com {len(solicitacoes_criadas)} item(ns) com sucesso!'
                        messages.success(request, mensagem)
                        return redirect('/solicitacoes/home/')
                        
                except Exception as e:
                    import traceback
                    error_trace = traceback.format_exc()
                    print(f"❌ ERRO ao criar solicitação de técnico: {e}")
                    print(f"📋 Traceback completo:\n{error_trace}")
                    
                    # Verificar se é requisição AJAX
                    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
                    
                    # Mensagem de erro mais detalhada
                    error_message = f'Erro ao criar solicitação de técnico: {str(e)}'
                    
                    # Se for um erro de integridade ou campo obrigatório, tentar extrair mensagem mais amigável
                    if 'IntegrityError' in str(type(e)):
                        error_message = 'Erro: Já existe uma solicitação com esse ticket. Tente usar um ID diferente.'
                    elif 'DoesNotExist' in str(type(e)):
                        error_message = f'Erro: Um dos registros referenciados não foi encontrado. Verifique se todos os dados estão corretos. Detalhes: {str(e)}'
                    elif 'ValueError' in str(type(e)):
                        error_message = f'Erro: Valor inválido em algum campo. Verifique os dados digitados. Detalhes: {str(e)}'
                    
                    if is_ajax:
                        return JsonResponse({
                            'success': False,
                            'message': error_message
                        }, status=400)
                    
                    messages.error(request, error_message)
                    return redirect('/solicitacoes/home/')
            
            # Verificar tipo de solicitação
            tipo_raw = request.POST.get('route', 'Casual')
            tipo = tipo_raw.lower().strip()
            if tipo == 'em rota' or tipo == 'em_rota':
                tipo = 'em_rota'
            else:
                tipo = 'casual'
            
            print(f"🔍 DEBUG: Tipo recebido do formulário: '{tipo_raw}' -> processado como: '{tipo}'")
            
            # Capturar dados comuns
            status = normalizar_status(request.POST.get('status') or 'pendente')
            if is_edit_mode:
                status = normalizar_status(solicitacao_existente.status)
            nome_do_recebedor = ''
            
            # Dados são capturados com prefixos diferentes baseado no tipo
            if tipo == 'em_rota':
                descricao = request.POST.get('route_description', '').strip()
                data_de_pagamento_str = request.POST.get('route_dataPagamento', '').strip()
                prioridade = request.POST.get('route_priority', 'baixa')
                anexo = request.FILES.get('route_anexos')
                # Para Em Rota, os valores detalhados estão apenas nos itens, não no geral
                # Apenas o valor de receita é geral
                valor_km = 0.0
                valor_pedagio = 0.0
                valor_hospedagem = 0.0
                valor_fluvial = 0.0
                valor_outros = 0.0
                valor_em_rota = limpar_valor_monetario(request.POST.get('route_valor_em_rota', '').strip())
                descricao_em_rota = request.POST.get('route_descricao_em_rota', '').strip()
                valor_receita = limpar_valor_monetario(request.POST.get('route_valor_receita', '').strip())
            else:  # Casual
                nome_do_recebedor = request.POST.get('casual_recebedor', '').strip()
                chave_pix_casual = request.POST.get('casual_pix', '').strip()
                cliente_empresa_casual = request.POST.get('casual_cliente_empresa', '').strip()
                cnpj_casual = request.POST.get('casual_cnpj', '').strip()
                descricao = request.POST.get('casual_description', '').strip()
                data_de_pagamento_str = request.POST.get('casual_dataPagamento', '').strip()
                prioridade = request.POST.get('casual_priority', 'baixa')
                anexo = request.FILES.get('casual_anexos')
                # Capturar valores detalhados para Casual
                valor_km_raw = request.POST.get('casual_valor_km', '').strip()
                valor_pedagio_raw = request.POST.get('casual_valor_pedagio', '').strip()
                valor_hospedagem_raw = request.POST.get('casual_valor_hospedagem', '').strip()
                valor_fluvial_raw = request.POST.get('casual_valor_fluvial', '').strip()
                valor_outros_raw = request.POST.get('casual_valor_outros', '').strip()
                valor_receita_raw = request.POST.get('casual_valor_receita', '').strip()
                
                valor_km = limpar_valor_monetario(valor_km_raw)
                valor_pedagio = limpar_valor_monetario(valor_pedagio_raw)
                valor_hospedagem = limpar_valor_monetario(valor_hospedagem_raw)
                valor_fluvial = limpar_valor_monetario(valor_fluvial_raw)
                valor_outros = limpar_valor_monetario(valor_outros_raw)
                valor_receita = limpar_valor_monetario(valor_receita_raw)
                
                if is_edit_mode:
                    print(f"🔍 DEBUG EDIÇÃO CASUAL - Valores capturados do formulário:")
                    print(f"   - casual_valor_km: '{valor_km_raw}' -> {valor_km}")
                    print(f"   - casual_valor_pedagio: '{valor_pedagio_raw}' -> {valor_pedagio}")
                    print(f"   - casual_valor_hospedagem: '{valor_hospedagem_raw}' -> {valor_hospedagem}")
                    print(f"   - casual_valor_fluvial: '{valor_fluvial_raw}' -> {valor_fluvial}")
                    print(f"   - casual_valor_outros: '{valor_outros_raw}' -> {valor_outros}")
                    print(f"   - casual_valor_receita: '{valor_receita_raw}' -> {valor_receita}")
                
                # Registrar recebedor no catálogo principal para reaproveitamento
                registrar_recebedor(nome_do_recebedor, chave_pix_casual)
            
            print(f"🔍 DEBUG Campos - Tipo: '{tipo}', Recebedor: '{nome_do_recebedor}', Descrição: '{descricao[:50]}...', Data: '{data_de_pagamento_str}', Prioridade: '{prioridade}'")
            
            # Validação básica dos campos obrigatórios (apenas para feedback do backend)
            # A validação principal deve ser feita no frontend
            # Se for requisição AJAX, retornar JSON ao invés de redirect
            is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
            
            if tipo == 'casual' and not nome_do_recebedor:
                if is_ajax:
                    return JsonResponse({
                        'success': False,
                        'message': 'O campo "Nome do Recebedor" é obrigatório.',
                        'field': 'casual_recebedor'
                    }, status=400)
                messages.error(request, 'O campo "Nome do Recebedor" é obrigatório.')
                return redirect('/solicitacoes/home/')
            
            if not descricao:
                if is_ajax:
                    field_name = 'casual_description' if tipo == 'casual' else 'route_description'
                    return JsonResponse({
                        'success': False,
                        'message': 'O campo "Descrição" é obrigatório.',
                        'field': field_name
                    }, status=400)
                messages.error(request, 'O campo "Descrição" é obrigatório.')
                return redirect('/solicitacoes/home/')
                
            if not data_de_pagamento_str:
                if is_ajax:
                    field_name = 'casual_dataPagamento' if tipo == 'casual' else 'route_dataPagamento'
                    return JsonResponse({
                        'success': False,
                        'message': 'O campo "Data de Pagamento" é obrigatório.',
                        'field': field_name
                    }, status=400)
                messages.error(request, 'O campo "Data de Pagamento" é obrigatório.')
                return redirect('/solicitacoes/home/')
            
            # Converter string de data para objeto Date
            try:
                data_de_pagamento = datetime.strptime(data_de_pagamento_str, '%Y-%m-%d').date()
            except (ValueError, TypeError):
                messages.error(request, 'Data de pagamento inválida. Use o formato correto.')
                return redirect('/solicitacoes/home/')
            
            # Arquivos já capturados acima baseado no tipo (route_anexos ou casual_anexos)
            
            # ✅ Calcular automaticamente tempo de criação (horário atual)
            tempo_criacao_auto = timezone.now().time()
            
            # ✅ Tempo na fila começa em 00:00:00 (será calculado dinamicamente no frontend)
            tempo_fila_inicial = time(0, 0, 0)
            
            if tipo == 'em_rota':
            
                # Processar solicitação Em Rota
                # Coletar todos os itens da rota (IDs dinâmicos)
                itens_rota = []
                valor_total = 0.0
                ticket_principal = None
                
                # Coletar todos os IDs dinamicamente (não apenas 1-4)
                # Verificar todos os campos route_id_N no POST
                route_ids_encontrados = []
                for key in request.POST.keys():
                    if key.startswith('route_id_'):
                        # Extrair o número do ID (ex: route_id_5 -> 5)
                        try:
                            num_id = int(key.replace('route_id_', ''))
                            route_ids_encontrados.append(num_id)
                        except ValueError:
                            continue
                
                # Ordenar para processar na ordem correta
                route_ids_encontrados.sort()
                print(f"🔍 DEBUG Em Rota - IDs encontrados: {route_ids_encontrados}")
                
                ordem_atual = 1
                recebedor_principal = ''
                chave_pix_principal = ''
                cliente_empresa_principal = ''
                cnpj_principal = ''
                for i in route_ids_encontrados:
                    route_id = request.POST.get(f'route_id_{i}', '').strip()
                    route_valor_raw = request.POST.get(f'route_valor_{i}', '').strip()
                    route_servico_id = request.POST.get(f'route_servico_{i}', '').strip()
                    route_recebedor = request.POST.get(f'route_recebedor_{i}', '').strip()
                    route_pix = request.POST.get(f'route_pix_{i}', '').strip()
                    route_cliente_empresa = request.POST.get(f'route_cliente_empresa_{i}', '').strip()
                    route_cnpj = request.POST.get(f'route_cnpj_{i}', '').strip()
                    
                    if route_id and not route_recebedor:
                        messages.error(request, f'Informe o recebedor para o ID {route_id}.')
                        return redirect('/solicitacoes/home/')
                    
                    if not recebedor_principal and route_recebedor:
                        recebedor_principal = route_recebedor
                    if not chave_pix_principal and route_pix:
                        chave_pix_principal = route_pix
                    if not cliente_empresa_principal and route_cliente_empresa:
                        cliente_empresa_principal = route_cliente_empresa
                    if not cnpj_principal and route_cnpj:
                        cnpj_principal = route_cnpj
                    
                    # Registrar recebedor para uso futuro
                    if route_recebedor:
                        registrar_recebedor(route_recebedor, route_pix)
                    
                    # Capturar valores detalhados para cada ID
                    valor_km = limpar_valor_monetario(request.POST.get(f'route_valor_km_{i}', '').strip())
                    valor_pedagio = limpar_valor_monetario(request.POST.get(f'route_valor_pedagio_{i}', '').strip())
                    valor_hospedagem = limpar_valor_monetario(request.POST.get(f'route_valor_hospedagem_{i}', '').strip())
                    valor_fluvial = limpar_valor_monetario(request.POST.get(f'route_valor_fluvial_{i}', '').strip())
                    valor_outros = limpar_valor_monetario(request.POST.get(f'route_valor_outros_{i}', '').strip())
                    
                    print(f"🔍 DEBUG Em Rota - Item {i} (ordem {ordem_atual}): ID='{route_id}', Valor='{route_valor_raw}', Servico='{route_servico_id}'")
                    
                    if route_id:  # Se há ID, deve ter serviço (valor é opcional)
                        # Validação: se ID está preenchido, serviço é obrigatório
                        if not route_servico_id:
                            messages.error(request, f'Para o ID {route_id}, o campo "Serviço" é obrigatório.')
                            return redirect('/solicitacoes/home/')
                        
                        # Limpar valor - mesmo tratamento do Casual (opcional)
                        valor_item = limpar_valor_monetario(route_valor_raw) if route_valor_raw else 0.0
                        print(f"🔍 DEBUG Em Rota - Item {i} valor processado: {valor_item}")
                        
                        # Se não há valor de atividade, calcular a partir dos valores detalhados
                        if valor_item <= 0:
                            valor_detalhados = valor_km + valor_pedagio + valor_hospedagem + valor_fluvial + valor_outros
                            if valor_detalhados > 0:
                                valor_item = valor_detalhados
                                print(f"🔍 DEBUG Em Rota - Item {i} valor calculado a partir dos detalhados: {valor_item}")
                            else:
                                # Se não há valor de atividade nem valores detalhados, permitir mas com valor 0
                                print(f"⚠️ DEBUG Em Rota - Item {i} sem valor de atividade nem valores detalhados")
                                valor_item = 0.0
                        
                        # Verificar se pelo menos um valor (atividade ou detalhados) foi preenchido
                        total_valores = valor_item + valor_km + valor_pedagio + valor_hospedagem + valor_fluvial + valor_outros
                        if total_valores <= 0:
                            messages.error(request, f'Para o ID {route_id}, preencha pelo menos um valor (Atividade ou algum custo adicional).')
                            return redirect('/solicitacoes/home/')
                        
                        # Buscar serviço
                        servico_obj = None
                        if route_servico_id:
                            try:
                                servico_obj = Servico.objects.get(id=route_servico_id, ativo=True)
                            except Servico.DoesNotExist:
                                messages.error(request, f'Serviço selecionado para o ID {route_id} não foi encontrado ou está inativo.')
                                return redirect('/solicitacoes/home/')
                        
                        if ordem_atual == 1:
                            ticket_principal = route_id
                        
                        itens_rota.append({
                            'ticket': route_id,
                            'valor': valor_item,
                            'servico': servico_obj,
                            'ordem': ordem_atual,
                            'recebedor': route_recebedor,
                            'chave_pix': route_pix,
                            'cliente_empresa': route_cliente_empresa,
                            'cnpj': route_cnpj,
                            'valor_km': valor_km,
                            'valor_pedagio': valor_pedagio,
                            'valor_hospedagem': valor_hospedagem,
                            'valor_fluvial': valor_fluvial,
                            'valor_outros': valor_outros
                        })
                        valor_total += valor_item
                        ordem_atual += 1
                
                if len(itens_rota) < 1:
                    messages.error(request, 'Solicitação Em Rota precisa de no mínimo 1 item preenchido.')
                    return redirect('/solicitacoes/home/')
                
                # Em solicitação agrupada (Em Rota), todos os itens devem ter o MESMO recebedor e chave PIX do primeiro item
                primeiro_recebedor = (itens_rota[0].get('recebedor') or '').strip()
                primeira_chave_pix = (itens_rota[0].get('chave_pix') or '').strip()
                for item in itens_rota:
                    item['recebedor'] = primeiro_recebedor
                    item['chave_pix'] = primeira_chave_pix
                nome_do_recebedor = recebedor_principal.strip() or primeiro_recebedor
                if not chave_pix_principal and primeira_chave_pix:
                    chave_pix_principal = primeira_chave_pix
                if not nome_do_recebedor:
                    messages.error(request, 'Informe pelo menos um recebedor nos itens da solicitação Em Rota.')
                    return redirect('/solicitacoes/home/')
                
                if not ticket_principal:
                    ticket_principal = itens_rota[0]['ticket']
                
                # ⚠️ VALIDAÇÃO: Verificar se algum ID da rota já existe no banco de dados
                tickets_rota = [item['ticket'] for item in itens_rota if item['ticket']]
                tickets_existentes = []
                for ticket in tickets_rota:
                    consulta_solic = Solicitacoes.objects.filter(ticket=ticket)
                    if is_edit_mode:
                        consulta_solic = consulta_solic.exclude(id=solicitacao_existente.id)
                    if consulta_solic.exists():
                        solicitacao_conflitante = consulta_solic.first()
                        if solicitacao_conflitante:
                            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                            tipo_existente = tipo_dict.get(solicitacao_conflitante.tipo, solicitacao_conflitante.tipo)
                            tickets_existentes.append(f"{ticket} (Solicitação {tipo_existente})")
                        else:
                            tickets_existentes.append(f"{ticket} (Solicitação)")
                        continue
                    
                    consulta_itens = SolicitacaoRotaItem.objects.filter(ticket_item=ticket)
                    if is_edit_mode:
                        consulta_itens = consulta_itens.exclude(solicitacao=solicitacao_existente)
                    if consulta_itens.exists():
                        item_existente = consulta_itens.first()
                        if item_existente and item_existente.solicitacao:
                            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                            tipo_existente = tipo_dict.get(item_existente.solicitacao.tipo, item_existente.solicitacao.tipo)
                            tickets_existentes.append(f"{ticket} (Item de rota {tipo_existente})")
                        else:
                            tickets_existentes.append(f"{ticket} (Item de rota)")
                
                if tickets_existentes:
                    tickets_str = ", ".join(tickets_existentes)
                    messages.error(
                        request, 
                        f'⚠️ Os seguintes IDs já estão cadastrados no sistema: {tickets_str}. '
                        f'Por favor, use IDs diferentes para os itens da rota.'
                    )
                    print(f"❌ ERRO: IDs já existem no banco de dados: {tickets_existentes}")
                    return redirect('/solicitacoes/home/')
                
                route_tickets = [item['ticket'] for item in itens_rota]
                titulo = f"Solicitação Em Rota - {', '.join(route_tickets)}"
                
                ticket_fallback = f"ROTA-{timezone.now().timestamp()}"
                if is_edit_mode and solicitacao_existente.ticket:
                    ticket_fallback = solicitacao_existente.ticket
                ticket_final = ticket_principal or ticket_fallback
                
                print(f"🔍 DEBUG Salvando solicitação Em Rota:")
                print(f"   - Ticket final: {ticket_final}")
                print(f"   - Valor Total: {valor_total}")
                print(f"   - Itens: {len(itens_rota)}")
                print(f"   - Modo edição: {is_edit_mode}")
                
                itens_para_criar = []
                for item in itens_rota:
                    itens_para_criar.append(SolicitacaoRotaItem(
                        solicitacao=None,  # Definido após obter a solicitação
                        ticket_item=item['ticket'],
                        valor=item['valor'],
                        servico=item['servico'],
                        ordem=item['ordem'],
                        recebedor=item.get('recebedor', ''),
                        chave_pix=item.get('chave_pix', ''),
                        cliente_empresa=item.get('cliente_empresa', ''),
                        cnpj=item.get('cnpj', ''),
                        valor_km=item.get('valor_km', 0.0),
                        valor_pedagio=item.get('valor_pedagio', 0.0),
                        valor_hospedagem=item.get('valor_hospedagem', 0.0),
                        valor_fluvial=item.get('valor_fluvial', 0.0),
                        valor_outros=item.get('valor_outros', 0.0)
                    ))
                
                with transaction.atomic():
                    if is_edit_mode:
                        solicitacao = solicitacao_existente
                        print(f"🔍 DEBUG EDIÇÃO EM ROTA - Valores antes do save:")
                        print(f"   - Valor Total: {valor_total} (anterior: {solicitacao.valor})")
                        print(f"   - Valor Receita: {valor_receita} (anterior: {solicitacao.valor_receita})")
                        print(f"   - Número de itens: {len(itens_rota)}")
                        solicitacao.ticket = ticket_final
                        solicitacao.titulo = titulo
                        solicitacao.nome_do_recebedor = nome_do_recebedor
                        solicitacao.valor = valor_total
                        solicitacao.descricao = descricao
                        solicitacao.data_de_pagamento = data_de_pagamento
                        solicitacao.prioridade = prioridade
                        solicitacao.servico = itens_rota[0]['servico']
                        solicitacao.tipo = 'em_rota'
                        solicitacao.valor_km = valor_km
                        solicitacao.valor_pedagio = valor_pedagio
                        solicitacao.valor_hospedagem = valor_hospedagem
                        solicitacao.valor_fluvial = valor_fluvial
                        solicitacao.valor_outros = valor_outros
                        solicitacao.valor_receita = valor_receita
                        solicitacao.valor_em_rota = valor_em_rota
                        solicitacao.descricao_em_rota = descricao_em_rota
                        solicitacao.chave_pix = chave_pix_principal or None
                        solicitacao.cliente_empresa = cliente_empresa_principal or None
                        solicitacao.cnpj = cnpj_principal or None
                        if anexo:
                            solicitacao.anexo = anexo
                        solicitacao.save()
                        print(f"✅ DEBUG EDIÇÃO EM ROTA - Solicitação salva com sucesso! ID: {solicitacao.id}")
                        solicitacao.itens_rota.all().delete()
                    else:
                        # Se criar com status aprovado ou concluido, preencher data_aprovacao
                        agora = timezone.now()
                        data_aprovacao_inicial = None
                        if status in ['aprovado', 'concluido']:
                            data_aprovacao_inicial = agora
                        
                        solicitacao = Solicitacoes.objects.create(
                            ticket=ticket_final,
                            status=status,
                            titulo=titulo,
                            nome_solicitante=request.user,
                            nome_do_recebedor=nome_do_recebedor,
                            chave_pix=chave_pix_principal or None,
                            cliente_empresa=cliente_empresa_principal or None,
                            cnpj=cnpj_principal or None,
                            valor=valor_total,
                            descricao=descricao,
                            data_de_pagamento=data_de_pagamento,
                            data_de_criacao=timezone.now().date(),
                            anexo=anexo,
                            tempo_criacao=tempo_criacao_auto,
                            tempo_fila=tempo_fila_inicial,
                            data_entrada_status=agora,
                            data_aprovacao=data_aprovacao_inicial,
                            prioridade=prioridade,
                            servico=itens_rota[0]['servico'],
                            tipo='em_rota',
                            valor_km=valor_km,
                            valor_pedagio=valor_pedagio,
                            valor_hospedagem=valor_hospedagem,
                            valor_fluvial=valor_fluvial,
                            valor_outros=valor_outros,
                            valor_receita=valor_receita,
                            valor_em_rota=valor_em_rota,
                            descricao_em_rota=descricao_em_rota
                        )
                    
                    for item in itens_para_criar:
                        item.solicitacao = solicitacao
                    SolicitacaoRotaItem.objects.bulk_create(itens_para_criar)
                
            else:
                # Processar solicitação Casual (original)
                id = request.POST.get('casual_id', '').strip()
                valor_raw = request.POST.get('casual_valor', '').strip()
                servico_id = request.POST.get('casual_service', '').strip()
                
                print(f"🔍 DEBUG Casual - ID='{id}', Valor='{valor_raw}', Servico='{servico_id}'")
                
                # ⚠️ VALIDAÇÃO: Verificar se o ID (ticket) já existe no banco de dados
                if id:
                    consulta_solic = Solicitacoes.objects.filter(ticket=id)
                    if is_edit_mode:
                        consulta_solic = consulta_solic.exclude(id=solicitacao_existente.id)
                    if consulta_solic.exists():
                        solicitacao_conf = consulta_solic.first()
                        if solicitacao_conf:
                            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                            tipo_existente = tipo_dict.get(solicitacao_conf.tipo, solicitacao_conf.tipo)
                            messages.error(
                                request, 
                                f'⚠️ O ID "{id}" já está cadastrado no sistema como solicitação "{tipo_existente}"! '
                                f'Por favor, use um ID diferente.'
                            )
                        else:
                            messages.error(
                                request, 
                                f'⚠️ O ID "{id}" já está cadastrado no sistema! '
                                f'Por favor, use um ID diferente.'
                            )
                        print(f"❌ ERRO: ID '{id}' já existe no banco de dados!")
                        return redirect('/solicitacoes/home/')
                    
                    consulta_item = SolicitacaoRotaItem.objects.filter(ticket_item=id)
                    if is_edit_mode:
                        consulta_item = consulta_item.exclude(solicitacao=solicitacao_existente)
                    if consulta_item.exists():
                        item_existente = consulta_item.first()
                        if item_existente and item_existente.solicitacao:
                            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                            tipo_existente = tipo_dict.get(item_existente.solicitacao.tipo, item_existente.solicitacao.tipo)
                            messages.error(
                                request, 
                                f'⚠️ O ID "{id}" já está cadastrado como item de rota em uma solicitação "{tipo_existente}"! '
                                f'Por favor, use um ID diferente.'
                            )
                        else:
                            messages.error(
                                request, 
                                f'⚠️ O ID "{id}" já está cadastrado como item de rota em outra solicitação! '
                                f'Por favor, use um ID diferente.'
                            )
                        print(f"❌ ERRO: ID '{id}' já existe como item de rota!")
                        return redirect('/solicitacoes/home/')
                
                # Limpar valor usando função auxiliar
                valor = limpar_valor_monetario(valor_raw)
                print(f"🔍 DEBUG Casual - Valor recebido do formulário: '{valor_raw}' -> processado: {valor}")
                
                # Se o valor for 0 ou não foi calculado, calcular a partir dos valores detalhados
                if valor == 0.0 or not valor_raw or valor_raw.strip() == '':
                    # Calcular valor total a partir dos valores detalhados + receita
                    valor_calculado = (
                        valor_km + valor_pedagio + valor_hospedagem + 
                        valor_fluvial + valor_outros + valor_receita
                    )
                    print(f"🔍 DEBUG Casual - Valores detalhados: KM={valor_km}, Pedágio={valor_pedagio}, Hospedagem={valor_hospedagem}, Fluvial={valor_fluvial}, Outros={valor_outros}, Receita={valor_receita}")
                    print(f"🔍 DEBUG Casual - Valor calculado a partir dos detalhados: {valor_calculado}")
                    
                    if valor_calculado > 0:
                        valor = valor_calculado
                        print(f"🔍 DEBUG - Valor recalculado a partir dos detalhados: {valor}")
                    else:
                        messages.error(request, 'O valor total da solicitação deve ser maior que zero. Preencha pelo menos um valor detalhado ou o valor de receita.')
                        print(f"❌ ERRO - Valor calculado é zero ou negativo: {valor_calculado}")
                        return redirect('/solicitacoes/home/')
                
                print(f"🔍 DEBUG - Valor final processado: {valor}")
                
                # Validação: valor deve ser maior que zero
                if valor <= 0:
                    messages.error(request, 'O valor total da solicitação deve ser maior que zero.')
                    print(f"❌ ERRO - Valor final é zero ou negativo: {valor}")
                    return redirect('/solicitacoes/home/')
                
                # Buscar o objeto Servico pelo ID
                servico_obj = None
                if servico_id:
                    try:
                        servico_obj = Servico.objects.get(id=servico_id, ativo=True)
                    except Servico.DoesNotExist:
                        messages.warning(request, f'Serviço selecionado não encontrado ou inativo.')
                
                # Gerar título e ticket
                titulo = f"Solicitação Casual - {id}" if id else "Solicitação Casual"
                ticket_fallback = f'CASUAL-{timezone.now().timestamp()}'
                if is_edit_mode and solicitacao_existente.ticket:
                    ticket_fallback = solicitacao_existente.ticket
                ticket_final = id if id else ticket_fallback
                
                print(f"🔍 DEBUG Salvando solicitação Casual:")
                print(f"   - ID/Ticket: {ticket_final}")
                print(f"   - Valor: {valor}")
                print(f"   - Serviço: {servico_obj}")
                print(f"   - Tipo: casual")
                print(f"   - Modo edição: {is_edit_mode}")
                
                try:
                    with transaction.atomic():
                        if is_edit_mode:
                            solicitacao = solicitacao_existente
                            print(f"🔍 DEBUG EDIÇÃO CASUAL - Valores antes do save:")
                            print(f"   - Valor: {valor} (anterior: {solicitacao.valor})")
                            print(f"   - Valor Receita: {valor_receita} (anterior: {solicitacao.valor_receita})")
                            print(f"   - Valor KM: {valor_km} (anterior: {solicitacao.valor_km})")
                            print(f"   - Recebedor: {nome_do_recebedor} (anterior: {solicitacao.nome_do_recebedor})")
                            solicitacao.ticket = ticket_final
                            solicitacao.titulo = titulo
                            solicitacao.nome_do_recebedor = nome_do_recebedor
                            solicitacao.chave_pix = chave_pix_casual or None
                            solicitacao.cliente_empresa = cliente_empresa_casual or None
                            solicitacao.cnpj = cnpj_casual or None
                            solicitacao.valor = valor
                            solicitacao.descricao = descricao
                            solicitacao.data_de_pagamento = data_de_pagamento
                            solicitacao.prioridade = prioridade
                            solicitacao.servico = servico_obj
                            solicitacao.tipo = 'casual'
                            solicitacao.valor_km = valor_km
                            solicitacao.valor_pedagio = valor_pedagio
                            solicitacao.valor_hospedagem = valor_hospedagem
                            solicitacao.valor_fluvial = valor_fluvial
                            solicitacao.valor_outros = valor_outros
                            solicitacao.valor_receita = valor_receita
                            solicitacao.valor_em_rota = 0.0
                            solicitacao.descricao_em_rota = ''
                            if anexo:
                                solicitacao.anexo = anexo
                            solicitacao.save()
                            print(f"✅ DEBUG EDIÇÃO CASUAL - Solicitação salva com sucesso! ID: {solicitacao.id}")
                            print(f"✅ DEBUG EDIÇÃO CASUAL - Valores após save: Valor={solicitacao.valor}, Receita={solicitacao.valor_receita}")
                            # Garantir que itens anteriores (se existirem) sejam removidos
                            solicitacao.itens_rota.all().delete()
                        else:
                            # Se criar com status aprovado ou concluido, preencher data_aprovacao
                            agora = timezone.now()
                            data_aprovacao_inicial = None
                            if status in ['aprovado', 'concluido']:
                                data_aprovacao_inicial = agora
                            
                            solicitacao = Solicitacoes.objects.create(
                                ticket=ticket_final,
                                status=status,
                                titulo=titulo,
                                nome_solicitante=request.user,
                                nome_do_recebedor=nome_do_recebedor,
                                chave_pix=chave_pix_casual or None,
                                cliente_empresa=cliente_empresa_casual or None,
                                cnpj=cnpj_casual or None,
                                valor=valor,
                                descricao=descricao,
                                data_de_pagamento=data_de_pagamento,
                                data_de_criacao=timezone.now().date(),
                                anexo=anexo,
                                tempo_criacao=tempo_criacao_auto,
                                tempo_fila=tempo_fila_inicial,
                                data_entrada_status=agora,
                                data_aprovacao=data_aprovacao_inicial,
                                prioridade=prioridade,
                                servico=servico_obj,
                                tipo='casual',
                                valor_km=valor_km,
                                valor_pedagio=valor_pedagio,
                                valor_hospedagem=valor_hospedagem,
                                valor_fluvial=valor_fluvial,
                                valor_outros=valor_outros,
                                valor_receita=valor_receita,
                                valor_em_rota=0.0,
                                descricao_em_rota=''
                            )
                    
                    print(f"✅ Solicitação Casual salva com sucesso! ID: {solicitacao.id}, Ticket: {solicitacao.ticket}")
                except Exception as e:
                    print(f"❌ ERRO ao salvar solicitação Casual: {e}")
                    messages.error(request, f'Erro ao salvar solicitação Casual: {str(e)}')
                    return redirect('/solicitacoes/home/')
            
            # Verificar se é requisição AJAX
            is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
            
            if is_ajax:
                # Retornar JSON para requisições AJAX
                return JsonResponse({
                    'success': True,
                    'message': success_message
                })
            
            messages.success(request, success_message)
            print(f"✅ DEBUG EDIÇÃO - Mensagem de sucesso enviada: {success_message}")
            if is_edit_mode:
                print(f"✅ DEBUG EDIÇÃO - Redirecionando após edição. Solicitação ID: {solicitacao_existente.id if solicitacao_existente else 'N/A'}")
            # Redirect para a mesma página para recarregar e mostrar o novo card
            return redirect('/solicitacoes/home/')
            
        except Exception as e:
            import traceback
            error_trace = traceback.format_exc()
            print(f"❌ ERRO GERAL ao processar solicitação: {str(e)}")
            print(f"📋 Traceback completo:\n{error_trace}")
            
            # Verificar se é requisição AJAX
            is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
            
            if is_ajax:
                return JsonResponse({
                    'success': False,
                    'message': f'Erro ao criar solicitação: {str(e)}'
                }, status=500)
            
            messages.error(request, f'Erro ao criar solicitação: {str(e)}')
            return redirect('/solicitacoes/home/')
    
    elif request.method == 'GET':
        """
        Verifica se o usuário que está dentro da requisição está autenticado, 
        se sim ele retorna e renderiza a página de home.
        Se não, retorna e redireciona para a área do login.
        """
        if request.user.is_authenticated:
            # Buscar as solicitações normalmente (migração já foi garantida no início da função)
            try:
                # Otimização: Buscar todas as solicitações de uma vez
                # e usar select_related e prefetch_related para otimizar queries
                # Retornar todas as solicitações (o filtro será aplicado no frontend)
                todas_solicitacoes = Solicitacoes.objects.select_related('nome_solicitante', 'servico').prefetch_related('itens_rota__servico', 'solicitacoes_tecnico').filter(excluida=False)
                
                solicitacoes_por_status = {
                    'pendente': [],
                    'recusado': [],
                    'aprovado': [],
                    'concluido': [],
                    'estorno': [],
                }
                
                objetos_para_corrigir = []
                for solicitacao in todas_solicitacoes:
                    status_normalizado = normalizar_status(solicitacao.status)
                    
                    # Ajustar o objeto em memória para manter consistência na renderização
                    if solicitacao.status != status_normalizado:
                        print(f"⚠️ Normalizando status da solicitação {solicitacao.id} ({solicitacao.ticket}): '{solicitacao.status}' -> '{status_normalizado}'")
                        solicitacao.status = status_normalizado
                        objetos_para_corrigir.append(solicitacao)
                    
                    solicitacoes_por_status.setdefault(status_normalizado, []).append(solicitacao)
                
                if objetos_para_corrigir:
                    Solicitacoes.objects.bulk_update(objetos_para_corrigir, ['status'])
                
                solicitacoes_pendentes = solicitacoes_por_status['pendente']
                solicitacoes_recusados = solicitacoes_por_status['recusado']
                solicitacoes_aprovado = solicitacoes_por_status['aprovado']
                solicitacoes_concluido = solicitacoes_por_status['concluido']
                solicitacoes_estorno = solicitacoes_por_status['estorno']
                
                return render(request, 'home/index.html', {
                    'solicitacoes_pendentes': solicitacoes_pendentes,
                    'num_solicitacoes_pendentes': len(solicitacoes_pendentes),

                    'solicitacoes_recusados': solicitacoes_recusados,
                    'num_solicitacoes_recusados': len(solicitacoes_recusados),

                    'solicitacoes_aprovado': solicitacoes_aprovado,
                    'num_solicitacoes_aprovado': len(solicitacoes_aprovado),

                    'solicitacoes_concluido': solicitacoes_concluido,
                    'num_solicitacoes_concluido': len(solicitacoes_concluido),
                    'solicitacoes_estorno': solicitacoes_estorno,
                    'num_solicitacoes_estorno': len(solicitacoes_estorno),
                })
            except Exception as e:
                import traceback
                error_trace = traceback.format_exc()
                print(f"❌ Erro ao buscar solicitações: {str(e)}")
                print(f"📋 Traceback:\n{error_trace}")
                # Retornar página vazia ao invés de erro
                return render(request, 'home/index.html', {
                    'solicitacoes_pendentes': [],
                    'num_solicitacoes_pendentes': 0,
                    'solicitacoes_recusados': [],
                    'num_solicitacoes_recusados': 0,
                    'solicitacoes_aprovado': [],
                    'num_solicitacoes_aprovado': 0,
                    'solicitacoes_concluido': [],
                    'num_solicitacoes_concluido': 0,
                    'solicitacoes_estorno': [],
                    'num_solicitacoes_estorno': 0,
                })
        else:
            return redirect('login')


@require_http_methods(["GET"])
def obter_itens_rota(request, solicitacao_id):
    """
    View para obter os itens da rota de uma solicitação via AJAX
    """
    try:
        solicitacao = Solicitacoes.objects.prefetch_related('itens_rota__servico').get(id=solicitacao_id)
        
        # Verificar se é uma solicitação "Em Rota"
        if solicitacao.tipo != 'em_rota':
            return JsonResponse({
                'success': False,
                'message': 'Esta solicitação não é do tipo "Em Rota"'
            }, status=400)
        
        # Buscar todos os itens da rota
        itens = solicitacao.itens_rota.all().order_by('ordem')
        
        # Serializar os itens com valores detalhados
        itens_data = []
        for item in itens:
            # Calcular valor da atividade para cada item
            # IMPORTANTE: item.valor já é o valor da atividade (não o valor total)
            valor_atividade_item = item.valor or 0.0
            
            itens_data.append({
                'ordem': item.ordem,
                'id': item.ticket_item,
                'valor': format_brl(item.valor),
                'servico': item.servico.nome if item.servico else 'N/A',
                'recebedor': item.recebedor or '',
                'chave_pix': item.chave_pix or '',
                'cliente_empresa': item.cliente_empresa or '',
                'cnpj': item.cnpj or '',
                'valor_km': format_brl(item.valor_km),
                'valor_pedagio': format_brl(item.valor_pedagio),
                'valor_hospedagem': format_brl(item.valor_hospedagem),
                'valor_fluvial': format_brl(item.valor_fluvial),
                'valor_outros': format_brl(item.valor_outros),
                'valor_total_item': format_brl(item.valor),
                'valor_atividade': format_brl(valor_atividade_item)  # Adicionar valor da atividade calculado
            })
        
        # Calcular valor total apenas com valores detalhados (sem atividade)
        valor_total_detalhados = solicitacao.get_valor_detalhados()
        
        # Calcular valor da receita como soma das atividades de todos os itens
        # IMPORTANTE: item.valor já é o valor da atividade (não o valor total)
        # Então basta somar todos os item.valor
        valor_receita_calculado = 0.0
        for item in itens:
            # item.valor já é o valor da atividade, não precisa subtrair nada
            valor_atividade_item = item.valor or 0.0
            valor_receita_calculado += valor_atividade_item
        
        return JsonResponse({
            'success': True,
            'itens': itens_data,
            'valor_total': format_brl(valor_total_detalhados),
            'valor_receita': format_brl(valor_receita_calculado),
            'valor_em_rota': format_brl(solicitacao.valor_em_rota),
            'descricao_em_rota': solicitacao.descricao_em_rota or ''
        })
        
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada'
        }, status=404)
        
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Erro ao obter itens da rota: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def obter_itens_tecnico(request, solicitacao_id):
    """
    View para obter os itens de técnico de uma solicitação via AJAX
    """
    try:
        solicitacao = Solicitacoes.objects.prefetch_related('solicitacoes_tecnico__servico', 'solicitacoes_tecnico__recebedor', 'solicitacoes_tecnico__cliente_empresa').get(id=solicitacao_id)
        
        # Verificar se é uma solicitação de técnico
        if not solicitacao.is_tecnico():
            return JsonResponse({
                'success': False,
                'message': 'Esta solicitação não é do tipo "Técnico"'
            }, status=400)
        
        # Buscar todos os itens de técnico
        itens_tecnico = solicitacao.solicitacoes_tecnico.all().order_by('id')
        
        # Serializar os itens
        itens_data = []
        ordem = 1
        for item in itens_tecnico:
            valor_total_item = (item.valor_pagamento_tecnico or 0.0) + (item.valor_extra or 0.0)
            
            # Usar o ticket_item se existir, senão usar o ticket da solicitação principal
            ticket_item = item.ticket_item if hasattr(item, 'ticket_item') and item.ticket_item else solicitacao.ticket
            
            itens_data.append({
                'ordem': ordem,
                'id': solicitacao.ticket,  # Usar o ticket da solicitação principal
                'ticket_tecnico': ticket_item,  # Ticket original do item
                'recebedor': item.recebedor.nome if item.recebedor else '',
                'recebedor_id': item.recebedor.id if item.recebedor else None,
                'chave_pix': item.recebedor.chave_pix if item.recebedor else '',
                'servico': item.servico.nome if item.servico else 'N/A',
                'servico_id': item.servico.id if item.servico else None,
                'cliente_empresa': item.cliente_empresa.nome if item.cliente_empresa else '',
                'cliente_empresa_id': item.cliente_empresa.id if item.cliente_empresa else None,
                'valor_pagamento': format_brl(item.valor_pagamento_tecnico),
                'valor_pagamento_raw': float(item.valor_pagamento_tecnico or 0),
                'valor_extra': format_brl(item.valor_extra) if item.valor_extra else 'R$ 0,00',
                'valor_extra_raw': float(item.valor_extra or 0),
                'valor': format_brl(valor_total_item),
                'valor_total_item': format_brl(valor_total_item),
                'descricao': item.descricao or '',
                'data_realizacao': item.data_realizacao_atividade.strftime('%Y-%m-%d') if item.data_realizacao_atividade else '',
                'data_pagamento': item.data_pagamento.strftime('%Y-%m-%d') if item.data_pagamento else '',
                'atividade_produtiva': 'Produtiva' if item.atividade_produtiva else 'Improdutiva',
                'atividade_produtiva_bool': item.atividade_produtiva
            })
            ordem += 1
        
        # Calcular valor total (soma de todos os itens)
        valor_total_calculado = sum([
            (item.valor_pagamento_tecnico or 0.0) + (item.valor_extra or 0.0)
            for item in itens_tecnico
        ])
        
        return JsonResponse({
            'success': True,
            'itens': itens_data,
            'valor_total': format_brl(valor_total_calculado),
        })
        
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada'
        }, status=404)
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao obter itens de técnico: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def obter_valores_detalhados_casual(request, solicitacao_id):
    """
    View para obter os valores detalhados de uma solicitação Casual via AJAX
    """
    try:
        solicitacao = Solicitacoes.objects.select_related('servico').get(id=solicitacao_id)
        
        # Verificar se é uma solicitação "Casual"
        if solicitacao.tipo != 'casual':
            return JsonResponse({
                'success': False,
                'message': 'Esta solicitação não é do tipo "Casual"'
            }, status=400)
        
        # Calcular valor total apenas com valores detalhados (sem atividade)
        valor_total_detalhados = solicitacao.get_valor_detalhados()
        
        # Calcular valor da receita como valor da atividade (valor_total - soma_detalhados)
        # O valor da receita é SOMENTE o valor da atividade, não o valor salvo no banco
        soma_detalhados_casual = (solicitacao.valor_km or 0.0) + (solicitacao.valor_pedagio or 0.0) + \
                                 (solicitacao.valor_hospedagem or 0.0) + (solicitacao.valor_fluvial or 0.0) + \
                                 (solicitacao.valor_outros or 0.0)
        valor_total_casual = solicitacao.valor or 0.0
        valor_receita_calculado = 0.0
        if valor_total_casual > soma_detalhados_casual:
            valor_receita_calculado = valor_total_casual - soma_detalhados_casual
        
        # Serializar os valores detalhados
        valores_detalhados = {
            'valor_km': format_brl(solicitacao.valor_km),
            'valor_pedagio': format_brl(solicitacao.valor_pedagio),
            'valor_hospedagem': format_brl(solicitacao.valor_hospedagem),
            'valor_fluvial': format_brl(solicitacao.valor_fluvial),
            'valor_outros': format_brl(solicitacao.valor_outros),
            'valor_receita': format_brl(valor_receita_calculado),  # Usar valor calculado, não o salvo
            'valor_total': format_brl(valor_total_detalhados),
            'servico': solicitacao.servico.nome if solicitacao.servico else 'N/A',
            'chave_pix': solicitacao.chave_pix or '',
            'cliente_empresa': solicitacao.cliente_empresa or '',
            'cnpj': solicitacao.cnpj or ''
        }
        
        return JsonResponse({
            'success': True,
            'valores': valores_detalhados
        })
        
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada'
        }, status=404)
        
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Erro ao obter valores detalhados: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def obter_detalhes_completos(request, solicitacao_id):
    """
    View para obter detalhes completos de uma solicitação (para relatório)
    Inclui todos os valores detalhados e informações completas
    """
    try:
        solicitacao = Solicitacoes.objects.select_related('servico', 'nome_solicitante').prefetch_related('itens_rota__servico', 'solicitacoes_tecnico').get(id=solicitacao_id)
        
        # Verificar se é solicitação de técnico
        is_tecnico = solicitacao.is_tecnico()
        
        ultimo_estorno = EstornoHistorico.objects.filter(
            solicitacao=solicitacao,
            status_novo='estorno'
        ).order_by('-data_hora').first()

        # Chave PIX: usar valor atual da tabela Recebedor (não o salvo na solicitação) para evitar exibir valor antigo no formulário
        chave_pix_casual = solicitacao.chave_pix or ''
        if solicitacao.tipo == 'casual' and not is_tecnico and solicitacao.nome_do_recebedor:
            receb = Recebedor.objects.filter(nome__iexact=solicitacao.nome_do_recebedor.strip()).first()
            if receb:
                chave_pix_casual = receb.chave_pix or ''
        # Informações básicas
        dados = {
            'id': solicitacao.id,
            'ticket': solicitacao.ticket,
            'titulo': solicitacao.titulo,
            'solicitante': str(solicitacao.nome_solicitante),
            'recebedor': solicitacao.nome_do_recebedor,
            'chave_pix': chave_pix_casual,
            'cliente_empresa': solicitacao.cliente_empresa or '',
            'cnpj': solicitacao.cnpj or '',
            'servico': solicitacao.servico.nome if solicitacao.servico else 'N/A',
            'servico_id': solicitacao.servico.id if solicitacao.servico else None,
            'tipo': solicitacao.tipo,
            'status': solicitacao.status,
            'prioridade': solicitacao.prioridade,
            'is_tecnico': is_tecnico,
            'justificativa_estorno': ultimo_estorno.justificativa if ultimo_estorno else '',
            'data_justificativa_estorno': ultimo_estorno.data_hora.strftime('%d/%m/%Y %H:%M') if ultimo_estorno else '',
            'descricao': solicitacao.descricao or 'N/A',
            'data_criacao': solicitacao.data_de_criacao.strftime('%d/%m/%Y') if solicitacao.data_de_criacao else 'N/A',
            'data_pagamento': solicitacao.data_de_pagamento.strftime('%d/%m/%Y') if solicitacao.data_de_pagamento else 'N/A',
            'data_criacao_iso': solicitacao.data_de_criacao.isoformat() if solicitacao.data_de_criacao else '',
            'data_pagamento_iso': solicitacao.data_de_pagamento.isoformat() if solicitacao.data_de_pagamento else '',
            'valor_total': format_brl(solicitacao.valor),
            'valor_total_raw': float(solicitacao.valor or 0),
            'valor_total_ticket': '',
            'valor_total_ticket_raw': 0.0,
            # Calcular valor_receita a partir das atividades (não usar o valor salvo)
            'valor_receita': '',  # Será calculado abaixo
            'valor_receita_raw': 0.0,  # Será calculado abaixo
            'valor_em_rota': format_brl(solicitacao.valor_em_rota),
            'valor_em_rota_raw': float(solicitacao.valor_em_rota or 0),
            'descricao_em_rota': solicitacao.descricao_em_rota or '',
            # Valores detalhados (sempre presentes)
            'valor_km': format_brl(solicitacao.valor_km),
            'valor_km_raw': float(solicitacao.valor_km or 0),
            'valor_pedagio': format_brl(solicitacao.valor_pedagio),
            'valor_pedagio_raw': float(solicitacao.valor_pedagio or 0),
            'valor_hospedagem': format_brl(solicitacao.valor_hospedagem),
            'valor_hospedagem_raw': float(solicitacao.valor_hospedagem or 0),
            'valor_fluvial': format_brl(solicitacao.valor_fluvial),
            'valor_fluvial_raw': float(solicitacao.valor_fluvial or 0),
            'valor_outros': format_brl(solicitacao.valor_outros),
            'valor_outros_raw': float(solicitacao.valor_outros or 0),
        }
        
        # Calcular valor_receita a partir das atividades (não usar o valor salvo no banco)
        # Se for "Em Rota", incluir itens da rota e calcular receita
        if solicitacao.tipo == 'em_rota':
            # Para Em Rota: soma das atividades de todos os itens
            # IMPORTANTE: item.valor já é o valor da atividade (não o valor total)
            valor_receita_calculado = 0.0
            itens_rota = solicitacao.itens_rota.all().order_by('ordem')
            for item in itens_rota:
                # item.valor já é o valor da atividade, não precisa subtrair nada
                valor_atividade_item = item.valor or 0.0
                valor_receita_calculado += valor_atividade_item
            dados['valor_receita'] = format_brl(valor_receita_calculado)
            dados['valor_receita_raw'] = float(valor_receita_calculado)
            # Mapa nome recebedor (lower) -> chave_pix atual (uma query para todos os itens, match case-insensitive)
            nomes_recebedores = list({(item.recebedor or '').strip() for item in itens_rota if (item.recebedor or '').strip()})
            recebedores_pix = {}
            if nomes_recebedores:
                if len(nomes_recebedores) == 1:
                    r = Recebedor.objects.filter(nome__iexact=nomes_recebedores[0]).only('nome', 'chave_pix').first()
                    if r:
                        recebedores_pix[r.nome.lower()] = r.chave_pix or ''
                else:
                    q = Q()
                    for n in nomes_recebedores:
                        q |= Q(nome__iexact=n)
                    for r in Recebedor.objects.filter(q).only('nome', 'chave_pix'):
                        recebedores_pix[r.nome.lower()] = r.chave_pix or ''
            dados['itens_rota'] = []
            valor_total_solicitacao_calculado = 0.0
            for item in itens_rota:
                nome_item = (item.recebedor or '').strip()
                chave_pix_item = (recebedores_pix.get(nome_item.lower(), item.chave_pix or '') if nome_item else (item.chave_pix or ''))
                valor_total_ticket_item = (
                    float(item.valor or 0)
                    + float(item.valor_km or 0)
                    + float(item.valor_pedagio or 0)
                    + float(item.valor_hospedagem or 0)
                    + float(item.valor_fluvial or 0)
                    + float(item.valor_outros or 0)
                )
                valor_total_solicitacao_calculado += valor_total_ticket_item
                dados['itens_rota'].append({
                    'ordem': item.ordem,
                    'ticket_item': item.ticket_item,
                    'servico': item.servico.nome if item.servico else 'N/A',
                    'servico_id': item.servico.id if item.servico else None,
                    'recebedor': item.recebedor or '',
                    'chave_pix': chave_pix_item,
                    'cliente_empresa': item.cliente_empresa or '',
                    'cnpj': item.cnpj or '',
                    'valor': format_brl(item.valor),
                    'valor_raw': float(item.valor or 0),
                    'valor_km': format_brl(item.valor_km),
                    'valor_km_raw': float(item.valor_km or 0),
                    'valor_pedagio': format_brl(item.valor_pedagio),
                    'valor_pedagio_raw': float(item.valor_pedagio or 0),
                    'valor_hospedagem': format_brl(item.valor_hospedagem),
                    'valor_hospedagem_raw': float(item.valor_hospedagem or 0),
                    'valor_fluvial': format_brl(item.valor_fluvial),
                    'valor_fluvial_raw': float(item.valor_fluvial or 0),
                    'valor_outros': format_brl(item.valor_outros),
                    'valor_outros_raw': float(item.valor_outros or 0),
                    'valor_total_ticket': format_brl(valor_total_ticket_item),
                    'valor_total_ticket_raw': valor_total_ticket_item,
                })
            # Na planilha, o total da solicitação agrupada deve ser a soma dos totais de cada ticket/item.
            dados['valor_total'] = format_brl(valor_total_solicitacao_calculado)
            dados['valor_total_raw'] = float(valor_total_solicitacao_calculado)
        else:
            # Para Casual: valor_total - soma_detalhados
            soma_detalhados_casual = (solicitacao.valor_km or 0.0) + (solicitacao.valor_pedagio or 0.0) + \
                                     (solicitacao.valor_hospedagem or 0.0) + (solicitacao.valor_fluvial or 0.0) + \
                                     (solicitacao.valor_outros or 0.0)
            valor_total_casual = solicitacao.valor or 0.0
            valor_receita_calculado = 0.0
            if valor_total_casual > soma_detalhados_casual:
                valor_receita_calculado = valor_total_casual - soma_detalhados_casual
            dados['valor_receita'] = format_brl(valor_receita_calculado)
            dados['valor_receita_raw'] = float(valor_receita_calculado)
            dados['valor_total_ticket'] = format_brl(valor_receita_calculado + soma_detalhados_casual)
            dados['valor_total_ticket_raw'] = float(valor_receita_calculado + soma_detalhados_casual)
        
        # Se for solicitação de técnico, adicionar itens de técnico
        if is_tecnico:
            itens_tecnico = solicitacao.solicitacoes_tecnico.all().order_by('id')
            dados['itens_tecnico'] = []
            clientes_principais = []
            cnpjs_principais = []
            descricoes_principais = []
            for item in itens_tecnico:
                valor_total_item = float(item.valor_pagamento_tecnico or 0) + float(item.valor_extra or 0)
                cliente_nome = item.cliente_empresa.nome if item.cliente_empresa else ''
                cliente_cnpj = item.cliente_empresa.cnpj if item.cliente_empresa and item.cliente_empresa.cnpj else ''
                descricao_item = item.descricao or ''
                if cliente_nome:
                    clientes_principais.append(cliente_nome)
                if cliente_cnpj:
                    cnpjs_principais.append(cliente_cnpj)
                if descricao_item:
                    descricoes_principais.append(descricao_item)
                dados['itens_tecnico'].append({
                    'ticket_tecnico': item.ticket_item or '',
                    'recebedor': item.recebedor.nome if item.recebedor else '',
                    'recebedor_id': item.recebedor.id if item.recebedor else None,
                    'chave_pix': item.recebedor.chave_pix if item.recebedor else '',
                    'servico': item.servico.nome if item.servico else '',
                    'servico_id': item.servico.id if item.servico else None,
                    'cliente_empresa': cliente_nome,
                    'cliente_empresa_id': item.cliente_empresa.id if item.cliente_empresa else None,
                    'cnpj': cliente_cnpj,
                    'valor_pagamento': format_brl(item.valor_pagamento_tecnico),
                    'valor_pagamento_raw': float(item.valor_pagamento_tecnico or 0),
                    'valor_extra': format_brl(item.valor_extra),
                    'valor_extra_raw': float(item.valor_extra or 0),
                    'valor': format_brl(valor_total_item),
                    'valor_raw': valor_total_item,
                    'valor_total_ticket': format_brl(valor_total_item),
                    'valor_total_ticket_raw': valor_total_item,
                    'descricao': descricao_item,
                    'data_realizacao': item.data_realizacao_atividade.strftime('%Y-%m-%d') if item.data_realizacao_atividade else '',
                    'data_pagamento': item.data_pagamento.strftime('%Y-%m-%d') if item.data_pagamento else '',
                    'atividade_produtiva': item.atividade_produtiva if hasattr(item, 'atividade_produtiva') else True,
                })
            # Em solicitação técnica, Cliente/CNPJ/Descrição costumam estar no item.
            # Espelha no nível principal para exportações que leem colunas principais.
            if not dados.get('cliente_empresa') and clientes_principais:
                dados['cliente_empresa'] = ' | '.join(dict.fromkeys(clientes_principais))
            if not dados.get('cnpj') and cnpjs_principais:
                dados['cnpj'] = ' | '.join(dict.fromkeys(cnpjs_principais))
            if not dados.get('descricao_em_rota') and descricoes_principais:
                dados['descricao_em_rota'] = ' | '.join(dict.fromkeys(descricoes_principais))
        
        return JsonResponse({
            'success': True,
            'dados': dados
        })
        
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada'
        }, status=404)
        
    except Exception as e:
        import traceback
        error_trace = traceback.format_exc()
        print(f"❌ Erro ao obter detalhes completos: {str(e)}")
        print(f"Traceback: {error_trace}")
        return JsonResponse({
            'success': False,
            'message': f'Erro ao obter detalhes: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def obter_contagens_solicitacoes(request):
    """
    View para obter contagens de solicitações por tipo (deslocamento e técnico)
    """
    try:
        # Contar solicitações de deslocamento (que não são de técnico)
        # Uma solicitação é de deslocamento se não tem nenhum item de técnico associado
        solicitacoes_deslocamento = Solicitacoes.objects.filter(
            excluida=False
        ).filter(
            ~Exists(SolicitacaoTecnico.objects.filter(solicitacao_id=OuterRef('pk')))
        )
        
        # Contar solicitações de técnico (que têm itens de técnico)
        solicitacoes_tecnico = Solicitacoes.objects.filter(
            excluida=False
        ).filter(
            Exists(SolicitacaoTecnico.objects.filter(solicitacao_id=OuterRef('pk')))
        )
        
        # Contar por status para deslocamento
        deslocamento_por_status = {
            'pendente': solicitacoes_deslocamento.filter(status='pendente').count(),
            'recusado': solicitacoes_deslocamento.filter(status='recusado').count(),
            'aprovado': solicitacoes_deslocamento.filter(status='aprovado').count(),
            'concluido': solicitacoes_deslocamento.filter(status='concluido').count(),
            'estorno': solicitacoes_deslocamento.filter(status='estorno').count(),
        }
        
        # Contar por status para técnico
        tecnico_por_status = {
            'pendente': solicitacoes_tecnico.filter(status='pendente').count(),
            'recusado': solicitacoes_tecnico.filter(status='recusado').count(),
            'aprovado': solicitacoes_tecnico.filter(status='aprovado').count(),
            'concluido': solicitacoes_tecnico.filter(status='concluido').count(),
            'estorno': solicitacoes_tecnico.filter(status='estorno').count(),
        }
        
        # Contar total de itens de técnico (SolicitacaoTecnico)
        total_itens_tecnico = SolicitacaoTecnico.objects.count()
        
        return JsonResponse({
            'success': True,
            'deslocamento': {
                'total': solicitacoes_deslocamento.count(),
                'por_status': deslocamento_por_status
            },
            'tecnico': {
                'total': solicitacoes_tecnico.count(),
                'total_itens': total_itens_tecnico,  # Total de itens na tabela SolicitacaoTecnico
                'por_status': tecnico_por_status
            }
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao obter contagens: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def exportar_relatorio_card(request, solicitacao_id):
    """
    View para exportar relatório detalhado de uma solicitação em nova guia
    """
    try:
        solicitacao = Solicitacoes.objects.prefetch_related(
            'itens_rota__servico',
            'solicitacoes_tecnico__recebedor',
            'solicitacoes_tecnico__servico',
            'solicitacoes_tecnico__cliente_empresa'
        ).get(id=solicitacao_id)
        
        ultimo_estorno = EstornoHistorico.objects.filter(
            solicitacao=solicitacao,
            status_novo='estorno'
        ).order_by('-data_hora').first()

        # Verificar se é solicitação de técnico
        is_tecnico = solicitacao.is_tecnico()
        
        # Calcular valores
        valor_total_detalhados = solicitacao.get_valor_detalhados()
        valor_receita_calculado = 0.0
        
        # Se for Em Rota, calcular receita e preparar itens
        itens_rota_data = []
        
        if is_tecnico:
            # Para solicitações de técnico, buscar todos os itens
            itens_tecnico = solicitacao.solicitacoes_tecnico.all().order_by('id')
            ordem = 1
            
            for item in itens_tecnico:
                valor_total_item = (item.valor_pagamento_tecnico or 0.0) + (item.valor_extra or 0.0)
                valor_receita_calculado += valor_total_item  # Para técnico, o valor total é a receita
                
                # Usar ticket_item se existir, senão usar o ticket da solicitação principal
                ticket_item = item.ticket_item if hasattr(item, 'ticket_item') and item.ticket_item else solicitacao.ticket
                
                # Buscar CNPJ do cliente/empresa se houver
                cnpj = ''
                if item.cliente_empresa and hasattr(item.cliente_empresa, 'cnpj'):
                    cnpj = item.cliente_empresa.cnpj or ''
                
                itens_rota_data.append({
                    'ordem': ordem,
                    'ticket_item': ticket_item,
                    'servico': item.servico.nome if item.servico else 'N/A',
                    'recebedor': item.recebedor.nome if item.recebedor else '',
                    'chave_pix': item.recebedor.chave_pix if item.recebedor else '',
                    'cliente_empresa': item.cliente_empresa.nome if item.cliente_empresa else '',
                    'cnpj': cnpj,
                    'valor_total': valor_total_item,
                    'valor_km': 0.0,  # Técnico não tem valores detalhados como deslocamento
                    'valor_pedagio': 0.0,
                    'valor_hospedagem': 0.0,
                    'valor_fluvial': 0.0,
                    'valor_outros': item.valor_extra or 0.0,  # valor_extra como "outros"
                    'soma_detalhados': item.valor_extra or 0.0,
                    'valor_atividade': item.valor_pagamento_tecnico or 0.0,  # valor principal
                    'data_realizacao': item.data_realizacao_atividade.strftime('%d/%m/%Y') if item.data_realizacao_atividade else '',
                    'data_pagamento': item.data_pagamento.strftime('%d/%m/%Y') if item.data_pagamento else '',
                    'atividade_produtiva': item.atividade_produtiva if hasattr(item, 'atividade_produtiva') else True,
                    'descricao': item.descricao or '',
                })
                ordem += 1
        elif solicitacao.tipo == 'em_rota':
            itens = solicitacao.itens_rota.all().order_by('ordem')
            for item in itens:
                # IMPORTANTE: item.valor já é o valor da atividade (não o valor total)
                valor_atividade_item = item.valor or 0.0
                valor_receita_calculado += valor_atividade_item
                
                # Calcular soma dos detalhados apenas para exibição
                soma_detalhados_item = (item.valor_km or 0.0) + (item.valor_pedagio or 0.0) + \
                                       (item.valor_hospedagem or 0.0) + (item.valor_fluvial or 0.0) + \
                                       (item.valor_outros or 0.0)
                valor_total_item = valor_atividade_item + soma_detalhados_item
                
                itens_rota_data.append({
                    'ordem': item.ordem,
                    'ticket_item': item.ticket_item,
                    'servico': item.servico.nome if item.servico else 'N/A',
                    'recebedor': item.recebedor or '',
                    'chave_pix': item.chave_pix or '',
                    'cliente_empresa': item.cliente_empresa or '',
                    'cnpj': item.cnpj or '',
                    'valor_total': valor_total_item,
                    'valor_km': item.valor_km or 0.0,
                    'valor_pedagio': item.valor_pedagio or 0.0,
                    'valor_hospedagem': item.valor_hospedagem or 0.0,
                    'valor_fluvial': item.valor_fluvial or 0.0,
                    'valor_outros': item.valor_outros or 0.0,
                    'soma_detalhados': soma_detalhados_item,
                    'valor_atividade': valor_atividade_item,
                    'atividade_produtiva': None,  # Em Rota não tem atividade produtiva/improdutiva
                })
        else:
            # Para Casual, calcular receita como valor da atividade (valor_total - soma_detalhados)
            # O valor da receita é SOMENTE o valor da atividade, não o valor salvo no banco
            soma_detalhados_casual = (solicitacao.valor_km or 0.0) + (solicitacao.valor_pedagio or 0.0) + \
                                     (solicitacao.valor_hospedagem or 0.0) + (solicitacao.valor_fluvial or 0.0) + \
                                     (solicitacao.valor_outros or 0.0)
            valor_total_casual = solicitacao.valor or 0.0
            if valor_total_casual > soma_detalhados_casual:
                valor_receita_calculado = valor_total_casual - soma_detalhados_casual
            
            # Para Casual, criar um item único na lista para exibir na tabela "Detalhamento por ID"
            itens_rota_data.append({
                'ordem': 1,
                'ticket_item': solicitacao.ticket,
                'servico': solicitacao.servico.nome if solicitacao.servico else 'N/A',
                'recebedor': solicitacao.nome_do_recebedor or '',
                'chave_pix': solicitacao.chave_pix or '',
                'cliente_empresa': solicitacao.cliente_empresa or '',
                'cnpj': solicitacao.cnpj or '',
                'valor_total': valor_total_casual,
                'valor_km': solicitacao.valor_km or 0.0,
                'valor_pedagio': solicitacao.valor_pedagio or 0.0,
                'valor_hospedagem': solicitacao.valor_hospedagem or 0.0,
                'valor_fluvial': solicitacao.valor_fluvial or 0.0,
                'valor_outros': solicitacao.valor_outros or 0.0,
                'soma_detalhados': soma_detalhados_casual,
                'valor_atividade': valor_receita_calculado,
                'atividade_produtiva': None,  # Casual não tem atividade produtiva/improdutiva
            })
        
        context = {
            'solicitacao': solicitacao,
            'valor_total_detalhados': valor_total_detalhados,
            'valor_receita': valor_receita_calculado,  # Usar valor calculado, não o salvo
            'itens_rota': itens_rota_data,
            'justificativa_estorno': ultimo_estorno.justificativa if ultimo_estorno else '',
            'data_justificativa_estorno': ultimo_estorno.data_hora if ultimo_estorno else None,
            'data_exportacao': timezone.now(),
        }
        
        return render(request, 'solicitacoes/relatorio_card.html', context)
        
    except Solicitacoes.DoesNotExist:
        messages.error(request, 'Solicitação não encontrada')
        return redirect('/solicitacoes/home/')
        
    except Exception as e:
        messages.error(request, f'Erro ao gerar relatório: {str(e)}')
        return redirect('/solicitacoes/home/')
 

@require_http_methods(["GET"])
def exportar_pix_ordens_pagamento(request):
    """
    Exporta ordens de pagamento com PIX (principalmente Técnico) em CSV ou XLSX.

    Query params:
      - format: csv | xlsx (default: xlsx)
      - include_rota: 1 para incluir itens Em Rota (opcional)
    """
    if not request.user.is_authenticated:
        return JsonResponse({'error': 'Não autenticado'}, status=401)
    if not user_can_view_dashboard(request.user):
        return JsonResponse({'error': 'Sem permissão'}, status=403)

    fmt = (request.GET.get('format') or 'xlsx').strip().lower()
    include_rota = (request.GET.get('include_rota') or '').strip() in ('1', 'true', 'True', 'yes', 'sim')

    headers = [
        'Ticket Solicitação',
        'Ordem',
        'Ticket Item',
        'Tipo',
        'Recebedor',
        'Chave PIX',
        'Serviço',
        'Cliente/Empresa',
        'CNPJ',
        'Valor Pagamento',
        'Valor Extra',
        'Valor Total',
        'Data Realização',
        'Data Pagamento',
        'Status',
        'Prioridade',
        'Descrição',
        'Atividade Produtiva',
    ]

    rows = []
    ordem_por_solicitacao = {}

    itens_tecnico = (
        SolicitacaoTecnico.objects.select_related('solicitacao', 'recebedor', 'servico', 'cliente_empresa')
        .all()
        .order_by('solicitacao__data_de_criacao', 'solicitacao_id', 'id')
    )
    for item in itens_tecnico:
        sol = item.solicitacao
        sol_id = sol.id if sol else None
        ordem_por_solicitacao[sol_id] = (ordem_por_solicitacao.get(sol_id, 0) + 1)
        ordem = ordem_por_solicitacao[sol_id]

        ticket_solic = sol.ticket if sol else ''
        ticket_item = (item.ticket_item or '').strip() or ticket_solic
        recebedor_nome = item.recebedor.nome if item.recebedor else ''
        chave_pix = item.recebedor.chave_pix if item.recebedor else ''
        servico_nome = item.servico.nome if item.servico else ''
        cliente_nome = item.cliente_empresa.nome if item.cliente_empresa else ''
        cnpj = (item.cliente_empresa.cnpj or '') if item.cliente_empresa and hasattr(item.cliente_empresa, 'cnpj') else ''
        valor_pag = float(item.valor_pagamento_tecnico or 0.0)
        valor_extra = float(item.valor_extra or 0.0)
        valor_total = valor_pag + valor_extra
        data_real = item.data_realizacao_atividade.isoformat() if item.data_realizacao_atividade else ''
        data_pag = item.data_pagamento.isoformat() if item.data_pagamento else ''
        status = sol.status if sol else ''
        prioridade = sol.prioridade if sol else ''
        descricao = item.descricao or ''
        atividade_prod = 'Sim' if getattr(item, 'atividade_produtiva', False) else 'Não'

        rows.append([
            ticket_solic,
            ordem,
            ticket_item,
            'Técnico',
            recebedor_nome,
            chave_pix,
            servico_nome,
            cliente_nome,
            cnpj,
            valor_pag,
            valor_extra,
            valor_total,
            data_real,
            data_pag,
            status,
            prioridade,
            descricao,
            atividade_prod,
        ])

    if include_rota:
        itens_rota = (
            SolicitacaoRotaItem.objects.select_related('solicitacao', 'servico', 'recebedor_fk')
            .all()
            .order_by('solicitacao__data_de_criacao', 'solicitacao_id', 'ordem', 'id')
        )
        for item in itens_rota:
            sol = item.solicitacao
            ticket_solic = sol.ticket if sol else ''
            ticket_item = (item.ticket_item or '').strip()
            recebedor_nome = (item.recebedor or '').strip() or (item.recebedor_fk.nome if item.recebedor_fk else '')
            chave_pix = (item.chave_pix or '').strip() or (item.recebedor_fk.chave_pix if item.recebedor_fk else '')
            servico_nome = item.servico.nome if item.servico else ''
            cliente_nome = (item.cliente_empresa or '').strip()
            cnpj = (item.cnpj or '').strip()
            valor_pag = float(item.valor or 0.0)
            valor_total = valor_pag + float(item.valor_km or 0.0) + float(item.valor_pedagio or 0.0) + float(item.valor_hospedagem or 0.0) + float(item.valor_fluvial or 0.0) + float(item.valor_outros or 0.0)
            data_pag = sol.data_de_pagamento.isoformat() if sol and sol.data_de_pagamento else ''
            status = sol.status if sol else ''
            prioridade = sol.prioridade if sol else ''

            rows.append([
                ticket_solic,
                int(item.ordem or 0),
                ticket_item,
                'Em Rota',
                recebedor_nome,
                chave_pix,
                servico_nome,
                cliente_nome,
                cnpj,
                valor_pag,
                0.0,
                float(valor_total),
                '',
                data_pag,
                status,
                prioridade,
                '',
                '',
            ])

    today = timezone.now().date().isoformat()
    if fmt == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename=\"pix_ordens_pagamento_{today}.csv\"'
        response.write('\ufeff')  # BOM UTF-8 para Excel
        writer = csv.writer(response, delimiter=';')
        writer.writerow(headers)
        for r in rows:
            writer.writerow(r)
        return response

    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
        except Exception as e:
            return JsonResponse({'error': f'openpyxl indisponível: {str(e)}'}, status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = 'PIX e Ordens de Pagamento'
        ws.append(headers)
        for r in rows:
            ws.append(r)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename=\"pix_ordens_pagamento_{today}.xlsx\"'
        return response

    return JsonResponse({'error': 'Formato inválido. Use format=csv ou format=xlsx.'}, status=400)


@require_http_methods(["POST"])
def exportar_solicitacoes_planilha(request):
    """
    Exporta planilha XLSX de solicitações (2 abas: deslocamento e técnico), com estilo fixo.
    Recebe JSON: {"ids": [1,2,3,...]}
    """
    try:
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except Exception:
        return JsonResponse({'error': 'JSON inválido.'}, status=400)

    ids = payload.get('ids') or []
    try:
        ids = [int(i) for i in ids]
    except Exception:
        return JsonResponse({'error': 'Lista de ids inválida.'}, status=400)
    ids = list(dict.fromkeys(ids))
    if not ids:
        return JsonResponse({'error': 'Nenhuma solicitação informada para exportação.'}, status=400)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except Exception as e:
        return JsonResponse({'error': f'openpyxl indisponível: {str(e)}'}, status=500)

    solicitacoes = (
        Solicitacoes.objects
        .filter(id__in=ids, excluida=False)
        .select_related('servico', 'nome_solicitante')
        .prefetch_related('itens_rota__servico', 'solicitacoes_tecnico__servico', 'solicitacoes_tecnico__recebedor', 'solicitacoes_tecnico__cliente_empresa')
    )
    solicitacoes_map = {s.id: s for s in solicitacoes}
    solicitacoes_ordenadas = [solicitacoes_map[i] for i in ids if i in solicitacoes_map]

    desloc_headers = [
        'ID', 'TIPO', 'SOLICITANTE', 'TICKET', 'SERVICO', 'STATUS', 'RECEBEDOR', 'CHAVE PIX',
        'CLIENTE/EMPRESA', 'CNPJ', 'VALOR DA ATIVIDADE', 'KM', 'PEDAGIO', 'HOSPEDAGEM',
        'FLUVIAL', 'OUTROS', 'VALOR TOTAL DO TICKET', 'VALOR TOTAL DA SOLICITAÇÃO',
        'DESCRICAO ITEM', 'CRIACAO', 'PAGAMENTO'
    ]
    tecnico_headers = [
        'ID', 'TIPO DE SOLICITAÇÃO', 'SOLICITANTE', 'TICKET', 'SERVICO', 'STATUS', 'RECEBEDOR', 'CHAVE PIX',
        'CLIENTE/EMPRESA', 'CNPJ', 'VALOR PAGAMENTO', 'VALOR EXTRA', 'VALOR TOTAL DO TICKET',
        'VALOR TOTAL DA SOLICITAÇÃO', 'DESCRICAO', 'DATA REALIZACAO', 'DATA PAGAMENTO', 'ATIVIDADE'
    ]

    def _usuario_nome(u):
        if not u:
            return ''
        full = (u.get_full_name() or '').strip()
        return full or (getattr(u, 'username', '') or '')

    def _fmt_num(n):
        return format_brl(float(n or 0))

    def _build_rows_for_solicitacao(sol):
        tipo = 'tecnico' if sol.is_tecnico() else (sol.tipo or '')
        solicitante = _usuario_nome(sol.nome_solicitante)
        status = sol.status or ''
        servico_principal = sol.servico.nome if sol.servico else ''
        data_criacao = sol.data_de_criacao.strftime('%d/%m/%Y') if sol.data_de_criacao else ''
        data_pagamento = sol.data_de_pagamento.strftime('%d/%m/%Y') if sol.data_de_pagamento else ''

        rows = []

        if sol.is_tecnico():
            itens = list(sol.solicitacoes_tecnico.all().order_by('id'))
            total_solic = 0.0
            for it in itens:
                valor_pagamento = float(it.valor_pagamento_tecnico or 0)
                valor_extra = float(it.valor_extra or 0)
                valor_total_ticket = valor_pagamento + valor_extra
                total_solic += valor_total_ticket
                rows.append([
                    sol.id, 'tecnico', solicitante, (it.ticket_item or sol.ticket or ''), (it.servico.nome if it.servico else servico_principal),
                    status, (it.recebedor.nome if it.recebedor else ''), (it.recebedor.chave_pix if it.recebedor else ''),
                    (it.cliente_empresa.nome if it.cliente_empresa else ''), (it.cliente_empresa.cnpj if it.cliente_empresa and it.cliente_empresa.cnpj else ''),
                    _fmt_num(valor_pagamento), _fmt_num(valor_extra), _fmt_num(valor_total_ticket),
                    '',
                    (it.descricao or ''),
                    (it.data_realizacao_atividade.strftime('%d/%m/%Y') if it.data_realizacao_atividade else ''),
                    (it.data_pagamento.strftime('%d/%m/%Y') if it.data_pagamento else data_pagamento),
                    ('Produtiva' if bool(it.atividade_produtiva) else 'Improdutiva')
                ])
            if rows:
                # VALOR TOTAL DA SOLICITAÇÃO na primeira linha do grupo; depois será mesclado por ID
                rows[0][13] = _fmt_num(total_solic)
            return rows

        if sol.tipo == 'em_rota':
            itens = list(sol.itens_rota.all().order_by('ordem'))
            total_solic = 0.0
            for it in itens:
                atividade = float(it.valor or 0)
                km = float(it.valor_km or 0)
                ped = float(it.valor_pedagio or 0)
                hosp = float(it.valor_hospedagem or 0)
                flu = float(it.valor_fluvial or 0)
                out = float(it.valor_outros or 0)
                total_ticket = atividade + km + ped + hosp + flu + out
                total_solic += total_ticket
                rows.append([
                    sol.id, 'em_rota', solicitante, (it.ticket_item or sol.ticket or ''), (it.servico.nome if it.servico else servico_principal),
                    status, (it.recebedor or ''), (it.chave_pix or ''), (it.cliente_empresa or ''), (it.cnpj or ''),
                    _fmt_num(atividade), _fmt_num(km), _fmt_num(ped), _fmt_num(hosp), _fmt_num(flu), _fmt_num(out),
                    _fmt_num(total_ticket), '', (sol.descricao_em_rota or ''), data_criacao, data_pagamento
                ])
            if rows:
                rows[0][17] = _fmt_num(total_solic)
            return rows

        # casual/deslocamento simples
        km = float(sol.valor_km or 0)
        ped = float(sol.valor_pedagio or 0)
        hosp = float(sol.valor_hospedagem or 0)
        flu = float(sol.valor_fluvial or 0)
        out = float(sol.valor_outros or 0)
        soma_detalhados = km + ped + hosp + flu + out
        total_casual = float(sol.valor or 0)
        atividade = total_casual - soma_detalhados if total_casual > soma_detalhados else 0.0
        total_ticket = atividade + soma_detalhados
        rows.append([
            sol.id, (sol.tipo or 'casual'), solicitante, (sol.ticket or ''), servico_principal, status,
            (sol.nome_do_recebedor or ''), (sol.chave_pix or ''), (sol.cliente_empresa or ''), (sol.cnpj or ''),
            _fmt_num(atividade), _fmt_num(km), _fmt_num(ped), _fmt_num(hosp), _fmt_num(flu), _fmt_num(out),
            _fmt_num(total_ticket), _fmt_num(total_ticket), (sol.descricao or ''), data_criacao, data_pagamento
        ])
        return rows

    desloc_rows = []
    tecnico_rows = []
    for sol in solicitacoes_ordenadas:
        rows = _build_rows_for_solicitacao(sol)
        if sol.is_tecnico():
            tecnico_rows.extend(rows)
        else:
            desloc_rows.extend(rows)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Solicitação deslocamento'
    ws2 = wb.create_sheet('Solicitação técnico')

    title_fill = PatternFill(fill_type='solid', fgColor='00B5E2')
    header_fill = PatternFill(fill_type='solid', fgColor='9EDFF0')
    thin_black = Side(style='thin', color='000000')
    border_all = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    title_font = Font(bold=True, size=15, color='000000')
    header_font = Font(bold=True, size=10, color='000000')
    body_font = Font(size=10, color='000000')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def _render_sheet(ws, title, rows, headers):
        ws.append([title] + [''] * (len(headers) - 1))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.append(headers)
        for r in rows:
            ws.append(r)

        ws.freeze_panes = 'A3'
        from openpyxl.utils import get_column_letter
        ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(2, ws.max_row)}"
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 20

        # Larguras para evitar quebra de linha no cabeçalho
        col_widths_map = {
            'ID': 8,
            'TIPO': 12,
            'TIPO DE SOLICITAÇÃO': 20,
            'SOLICITANTE': 24,
            'TICKET': 14,
            'SERVICO': 24,
            'STATUS': 12,
            'RECEBEDOR': 18,
            'CHAVE PIX': 18,
            'CLIENTE/EMPRESA': 22,
            'CNPJ': 16,
            'VALOR DA ATIVIDADE': 18,
            'KM': 10,
            'PEDAGIO': 12,
            'HOSPEDAGEM': 14,
            'FLUVIAL': 10,
            'OUTROS': 10,
            'VALOR TOTAL DO TICKET': 22,
            'VALOR TOTAL DA SOLICITAÇÃO': 28,
            'DESCRICAO ITEM': 26,
            'CRIACAO': 12,
            'PAGAMENTO': 12,
            'VALOR PAGAMENTO': 18,
            'VALOR EXTRA': 14,
            'VALOR TOTAL': 16,
            'VALOR TOTAL DO TICKET': 22,
            'VALOR TOTAL DA SOLICITAÇÃO': 28,
            'DESCRICAO': 24,
            'DATA REALIZACAO': 16,
            'DATA PAGAMENTO': 16,
            'ATIVIDADE': 14,
        }
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = col_widths_map.get(headers[col - 1], 14)

        idx_id = 1
        idx_total_sol = headers.index('VALOR TOTAL DA SOLICITAÇÃO') + 1 if 'VALOR TOTAL DA SOLICITAÇÃO' in headers else None

        # Mesclar por ID e centralizar ID + total solicitação
        if rows and idx_total_sol is not None:
            start = 3
            while start <= ws.max_row:
                current_id = ws.cell(start, idx_id).value
                end = start + 1
                while end <= ws.max_row and ws.cell(end, idx_id).value == current_id:
                    end += 1
                block_end = end - 1
                if block_end > start:
                    ws.merge_cells(start_row=start, start_column=idx_id, end_row=block_end, end_column=idx_id)
                    ws.merge_cells(start_row=start, start_column=idx_total_sol, end_row=block_end, end_column=idx_total_sol)
                start = end

        for r in range(1, ws.max_row + 1):
            for c in range(1, len(headers) + 1):
                cell = ws.cell(r, c)
                cell.border = border_all
                if r == 1:
                    cell.fill = title_fill
                    cell.font = title_font
                    cell.alignment = center
                elif r == 2:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False, shrink_to_fit=True)
                else:
                    cell.font = body_font
                    if c == idx_id or (idx_total_sol is not None and c == idx_total_sol):
                        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

    _render_sheet(ws1, 'SOLICITAÇÃO DE DESLOCAMENTO', desloc_rows, desloc_headers)
    _render_sheet(ws2, 'SOLICITAÇÃO DE TÉCNICO', tecnico_rows, tecnico_headers)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    today = timezone.now().date().isoformat()
    response = HttpResponse(
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=\"solicitacoes_geral_{today}.xlsx\"'
    return response


@group_required('Administrador', 'Financeiro')
@require_http_methods(["GET"])
def auditoria_logs(request):
    """
    Página de auditoria para visualizar logs CRUD.
    """
    logs = AuditoriaLog.objects.select_related('usuario').all()

    q = (request.GET.get('q') or '').strip()
    acao = (request.GET.get('acao') or '').strip()
    app = (request.GET.get('app') or '').strip()
    modelo = (request.GET.get('modelo') or '').strip()
    show_auto = (request.GET.get('show_auto') or '').strip() == '1'

    if q:
        # Busca sem expor IP (dado sensível; continua gravado no banco)
        logs = logs.filter(
            Q(objeto_id__icontains=q) |
            Q(usuario__username__icontains=q) |
            Q(origem__icontains=q)
        )
    if acao:
        logs = logs.filter(acao=acao)
    if app:
        logs = logs.filter(app=app)
    if modelo:
        logs = logs.filter(modelo=modelo)

    # Por padrão, oculta deletes automáticos de itens internos durante edição de solicitação.
    # Se o usuário quiser ver tudo, basta marcar "Exibir ações automáticas".
    if not show_auto and not modelo:
        logs = logs.exclude(
            acao='delete',
            modelo__in=['solicitacaorotaitem', 'solicitacaotecnico'],
            origem__startswith='/solicitacoes/home/',
        )

    paginator = Paginator(logs.order_by('-data_hora'), 500)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Labels em português para exibição dos campos na auditoria
    AUDIT_FIELD_LABELS = {
        'chave_pix': 'Chave PIX',
        'nome': 'Nome',
        'nome_do_recebedor': 'Nome do Recebedor',
        'recebedor': 'Recebedor',
        'ticket': 'Ticket',
        'titulo': 'Título',
        'status': 'Status',
        'prioridade': 'Prioridade',
        'valor': 'Valor',
        'valor_receita': 'Valor Receita',
        'valor_km': 'Valor KM',
        'valor_pedagio': 'Valor Pedágio',
        'valor_hospedagem': 'Valor Hospedagem',
        'valor_fluvial': 'Valor Fluvial',
        'valor_outros': 'Valor Outros',
        'cliente_empresa': 'Cliente/Empresa',
        'cnpj': 'CNPJ',
        'servico': 'Serviço',
        'descricao': 'Descrição',
        'data_de_pagamento': 'Data de Pagamento',
        'data_de_criacao': 'Data de Criação',
        'valor_pagamento_tecnico': 'Valor Pagamento Técnico',
        'valor_extra': 'Valor Extra',
        'ativo': 'Ativo',
        'data_criacao': 'Data de Criação',
        'data_atualizacao': 'Data de Atualização',
    }

    # Resolver ticket da solicitação para cada log que tem solicitacao_id
    ids_solic = list({log.solicitacao_id for log in page_obj if getattr(log, 'solicitacao_id', None) is not None})
    tickets_por_id = {}
    if ids_solic:
        for sid, ticket in Solicitacoes.objects.filter(pk__in=ids_solic).values_list('id', 'ticket'):
            tickets_por_id[sid] = ticket or '-'

    for log in page_obj:
        log.ticket_solicitacao = tickets_por_id.get(getattr(log, 'solicitacao_id', None), None) if getattr(log, 'solicitacao_id', None) else None
        cambios = []
        antes = log.dados_anteriores or {}
        depois = log.dados_novos or {}
        campos = log.campos_alterados or []
        if not campos:
            campos = list(set(list(antes.keys()) + list(depois.keys())))
        for campo in campos:
            val_antes = antes.get(campo, '-')
            val_depois = depois.get(campo, '-')
            if isinstance(val_antes, (dict, list)): val_antes = json.dumps(val_antes, ensure_ascii=False)
            if isinstance(val_depois, (dict, list)): val_depois = json.dumps(val_depois, ensure_ascii=False)
            label = AUDIT_FIELD_LABELS.get(campo, campo.replace('_', ' ').title())
            cambios.append({
                'campo': campo,
                'label': label,
                'antes': val_antes,
                'depois': val_depois,
            })
        log.detalhes_formatados = cambios
        log.detalhes_json = json.dumps(cambios, ensure_ascii=False)
        log.campos_alterados_display = ', '.join([AUDIT_FIELD_LABELS.get(c, c.replace('_', ' ').title()) for c in (log.campos_alterados or [])]) or '-'

    context = {
        'page_obj': page_obj,
        'q': q,
        'acao': acao,
        'app': app,
        'modelo': modelo,
        'show_auto': show_auto,
        'acoes_disponiveis': AuditoriaLog.objects.values_list('acao', flat=True).distinct(),
        'apps_disponiveis': AuditoriaLog.objects.values_list('app', flat=True).distinct(),
        'modelos_disponiveis': AuditoriaLog.objects.values_list('modelo', flat=True).distinct(),
    }
    return render(request, 'auditoria/index.html', context)


@require_http_methods(["POST"])
def atualizar_status(request):
    """
    View para atualizar o status de uma solicitação via AJAX
    Apenas Financeiro e Administrador podem mudar status
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'success': False,
            'message': 'Não autenticado'
        }, status=401)
    
    # Verificar permissão: apenas Financeiro e Admin podem mudar status
    from usuarios.decorators import user_can_change_status, user_is_financeiro
    if not user_can_change_status(request.user):
        return JsonResponse({
            'success': False,
            'message': 'Você não tem permissão para alterar o status de solicitações.'
        }, status=403)
    
    try:
        # Ler dados do corpo da requisição
        data = json.loads(request.body)
        card_id = data.get('card_id')
        new_status = data.get('status')
        justificativa_estorno = (data.get('justificativa_estorno') or '').strip()
        
        # Mapear filas do Kanban para status do banco
        status_mapping = {
            'planning': 'pendente',
            'test': 'recusado',
            'launch': 'aprovado',
            'success': 'concluido',
            'refund': 'estorno',
        }
        
        # Validar dados
        if not card_id or not new_status:
            return JsonResponse({
                'success': False,
                'message': 'Dados inválidos'
            }, status=400)
        
        # Converter fila para status
        status_db = normalizar_status(status_mapping.get(new_status, new_status))

        # Apenas Financeiro pode enviar para ESTORNO
        if status_db == 'estorno' and not user_is_financeiro(request.user):
            return JsonResponse({
                'success': False,
                'message': 'Apenas usuários do grupo Financeiro podem mover solicitações para Estorno.'
            }, status=403)
        
        # Buscar e atualizar solicitação
        from django.utils import timezone
        solicitacao = Solicitacoes.objects.get(id=card_id, excluida=False)
        status_anterior = solicitacao.status
        agora = timezone.now()
        
        # Se o status mudou, atualizar data_entrada_status
        if status_anterior != status_db:
            solicitacao.data_entrada_status = agora

        # Justificativa obrigatória ao mover para Estorno
        if status_db == 'estorno' and status_anterior != 'estorno' and not justificativa_estorno:
            return JsonResponse({
                'success': False,
                'message': 'Informe a justificativa para mover a solicitação para Estorno.'
            }, status=400)
        
        # Registrar momento exato da aprovação para manter a métrica mesmo após outras mudanças
        if status_db == 'aprovado' and not solicitacao.data_aprovacao:
            solicitacao.data_aprovacao = agora
        elif status_db == 'concluido' and not solicitacao.data_aprovacao:
            # Alguns fluxos podem pular direto para concluído; registramos a aprovação no mesmo instante
            solicitacao.data_aprovacao = agora
        
        solicitacao.status = status_db
        solicitacao.save()

        # Registrar histórico específico de estorno em tabela dedicada
        if status_db == 'estorno' and status_anterior != 'estorno':
            EstornoHistorico.objects.create(
                solicitacao=solicitacao,
                usuario=request.user,
                status_anterior=status_anterior,
                status_novo='estorno',
                justificativa=justificativa_estorno,
                origem=request.path,
            )

        # Garantir registro de auditoria para toda movimentação de card
        garantir_log_movimentacao_card(
            request=request,
            solicitacao=solicitacao,
            status_anterior=status_anterior,
            status_novo=status_db,
            justificativa_estorno=justificativa_estorno,
        )
        
        # Buscar contadores atualizados do banco de dados
        contadores = {
            'pendente': Solicitacoes.objects.filter(status='pendente', excluida=False).count(),
            'recusado': Solicitacoes.objects.filter(status='recusado', excluida=False).count(),
            'aprovado': Solicitacoes.objects.filter(status='aprovado', excluida=False).count(),
            'concluido': Solicitacoes.objects.filter(status='concluido', excluida=False).count(),
            'estorno': Solicitacoes.objects.filter(status='estorno', excluida=False).count(),
        }
        
        return JsonResponse({
            'success': True,
            'message': f'Status atualizado para {status_db}',
            'card_id': card_id,
            'new_status': status_db,
            'counters': contadores
        })
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada'
        }, status=404)
    except json.JSONDecodeError:
        return JsonResponse({
            'success': False,
            'message': 'Erro ao processar dados JSON'
        }, status=400)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Erro ao atualizar status: {str(e)}'
        }, status=500)

@require_http_methods(["GET"])
def buscar_recebedores(request):
    """
    View para buscar recebedores via AJAX
    Retorna lista de recebedores ativos para preencher datalist
    Pode buscar por ID específico ou por nome
    """
    try:
        query = request.GET.get('q', '').strip()
        recebedor_id = request.GET.get('id', '').strip()
        
        # Se há ID específico, buscar apenas esse recebedor
        if recebedor_id:
            try:
                recebedor = Recebedor.objects.get(id=int(recebedor_id))
                recebedores_list = [{
                    'id': recebedor.id,
                    'nome': recebedor.nome,
                    'chave_pix': recebedor.chave_pix,
                    'supervisor': recebedor.supervisor or '',
                    'ativo': recebedor.ativo
                }]
            except (Recebedor.DoesNotExist, ValueError):
                recebedores_list = []
        else:
            recebedores = Recebedor.objects.all()
            
            if query:
                # Buscar por nome (inclui inativos se houver query)
                recebedores = recebedores.filter(nome__icontains=query)
            else:
                # Se não há query, retornar apenas ativos ordenados por nome
                recebedores = recebedores.filter(ativo=True).order_by('nome')
            
            # Se não há query ou query vazia, retornar todos os ativos (sem limite)
            if not query:
                limit = None
            else:
                limit = 20  # Limitar apenas para buscas
            
            recebedores_queryset = recebedores[:limit] if limit else recebedores
            
            recebedores_list = [
                {
                    'id': r.id,
                    'nome': r.nome,
                    'chave_pix': r.chave_pix,
                    'supervisor': r.supervisor or '',
                    'ativo': r.ativo
                }
                for r in recebedores_queryset
            ]
        
        return JsonResponse({
            'success': True,
            'recebedores': recebedores_list
        })
    except Exception as e:
        print(f"❌ Erro ao buscar recebedores: {e}")
        return JsonResponse({
            'success': False,
            'recebedores': [],
            'message': str(e)
        }, status=500)

@require_http_methods(["GET"])
def buscar_clientes_empresas(request):
    """
    View para buscar clientes/empresas via AJAX
    Retorna lista de clientes/empresas ativos para preencher datalist
    Pode buscar por ID específico ou por nome
    """
    try:
        query = request.GET.get('q', '').strip()
        cliente_id = request.GET.get('id', '').strip()
        
        # Se há ID específico, buscar apenas esse cliente
        if cliente_id:
            try:
                cliente = ClienteEmpresa.objects.get(id=int(cliente_id))
                clientes_list = [{
                    'id': cliente.id,
                    'nome': cliente.nome,
                    'cnpj': cliente.cnpj or '',
                    'ativo': cliente.ativo
                }]
            except (ClienteEmpresa.DoesNotExist, ValueError):
                clientes_list = []
        else:
            clientes = ClienteEmpresa.objects.all()
            
            if query:
                # Buscar por nome (inclui inativos se houver query)
                clientes = clientes.filter(nome__icontains=query)
            else:
                # Se não há query, retornar apenas ativos
                clientes = clientes.filter(ativo=True)
            
            clientes_list = [
                {
                    'id': c.id,
                    'nome': c.nome,
                    'cnpj': c.cnpj or '',
                    'ativo': c.ativo
                }
                for c in clientes[:20]  # Limitar a 20 resultados
            ]
        
        return JsonResponse({
            'success': True,
            'clientes_empresas': clientes_list
        })
    except Exception as e:
        print(f"❌ Erro ao buscar clientes/empresas: {e}")
        return JsonResponse({
            'success': False,
            'clientes_empresas': [],
            'message': str(e)
        }, status=500)
@require_http_methods(["POST"])
def excluir_solicitacao(request, solicitacao_id):
    """
    View para excluir uma solicitação individual
    Apenas usuários do grupo Financeiro podem executar esta ação
    E apenas para solicitações recusadas
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'success': False,
            'message': 'Não autenticado'
        }, status=401)
    
    # Verificar permissão: apenas Financeiro pode excluir solicitações
    from usuarios.decorators import user_is_financeiro
    if not user_is_financeiro(request.user):
        return JsonResponse({
            'success': False,
            'message': 'Apenas membros do grupo Financeiro podem excluir solicitações.'
        }, status=403)
    
    try:
        # Buscar a solicitação não excluída
        solicitacao = Solicitacoes.objects.get(id=solicitacao_id, excluida=False)
        
        # Verificar se é recusada
        if solicitacao.status != 'recusado':
            return JsonResponse({
                'success': False,
                'message': 'Apenas solicitações recusadas podem ser excluídas.'
            }, status=400)
        
        ticket = solicitacao.ticket
        registrar_exclusao_solicitacao(
            solicitacao=solicitacao,
            usuario=request.user,
            origem=request.path,
            motivo="Exclusão individual de solicitação recusada",
        )
        solicitacao.excluida = True
        solicitacao.data_exclusao = timezone.now()
        solicitacao.excluida_por = request.user
        solicitacao.save(update_fields=['excluida', 'data_exclusao', 'excluida_por'])
        
        return JsonResponse({
            'success': True,
            'message': f'Solicitação #{ticket} excluída com sucesso!'
        })
        
    except Solicitacoes.DoesNotExist:
        return JsonResponse({
            'success': False,
            'message': 'Solicitação não encontrada.'
        }, status=404)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao excluir solicitação: {str(e)}'
        }, status=500)

@require_http_methods(["POST"])
def excluir_solicitacoes_recusadas(request):
    """
    View para excluir todas as solicitações recusadas
    Apenas usuários do grupo Financeiro podem executar esta ação
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'success': False,
            'message': 'Não autenticado'
        }, status=401)
    
    # Verificar permissão: apenas Financeiro pode excluir solicitações recusadas
    from usuarios.decorators import user_is_financeiro
    if not user_is_financeiro(request.user):
        return JsonResponse({
            'success': False,
            'message': 'Apenas membros do grupo Financeiro podem excluir solicitações recusadas.'
        }, status=403)
    
    try:
        # Buscar todas as solicitações recusadas não excluídas
        solicitacoes_recusadas = Solicitacoes.objects.filter(status='recusado', excluida=False)
        quantidade = solicitacoes_recusadas.count()
        
        if quantidade == 0:
            return JsonResponse({
                'success': True,
                'message': 'Não há solicitações recusadas para excluir.',
                'quantidade': 0
            })
        
        # Exclusão lógica em lote
        agora = timezone.now()
        for solicitacao in solicitacoes_recusadas:
            registrar_exclusao_solicitacao(
                solicitacao=solicitacao,
                usuario=request.user,
                origem=request.path,
                motivo="Exclusão em lote de solicitações recusadas",
            )
            solicitacao.excluida = True
            solicitacao.data_exclusao = agora
            solicitacao.excluida_por = request.user
            solicitacao.save(update_fields=['excluida', 'data_exclusao', 'excluida_por'])
        
        return JsonResponse({
            'success': True,
            'message': f'{quantidade} solicitação(ões) recusada(s) excluída(s) com sucesso!',
            'quantidade': quantidade
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao excluir solicitações: {str(e)}'
        }, status=500)


@require_http_methods(["GET"])
def verificar_id_existente(request):
    """
    View para verificar se um ID (ticket) já está cadastrado no sistema.
    Usado para validação em tempo real no formulário.
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'existe': False,
            'message': 'Não autenticado'
        }, status=401)
    
    ticket = request.GET.get('ticket', '').strip()
    solicitacao_id = request.GET.get('solicitacao_id', '').strip()  # Para modo edição
    
    if not ticket:
        return JsonResponse({
            'existe': False,
            'message': 'ID não fornecido'
        })
    
    try:
        # Verificar se existe como ticket principal em Solicitacoes
        consulta_solic = Solicitacoes.objects.filter(ticket=ticket)
        if solicitacao_id:
            try:
                consulta_solic = consulta_solic.exclude(id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_solic.exists():
            solicitacao_existente = consulta_solic.first()
            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
            tipo_existente = tipo_dict.get(solicitacao_existente.tipo, solicitacao_existente.tipo)
            status_existente = solicitacao_existente.status
            
            return JsonResponse({
                'existe': True,
                'tipo': 'solicitacao',
                'message': f'O ID "{ticket}" já está cadastrado como solicitação "{tipo_existente}" (Status: {status_existente}).',
                'detalhes': {
                    'tipo': tipo_existente,
                    'status': status_existente,
                    'titulo': solicitacao_existente.titulo
                }
            })
        
        # Verificar se existe como ticket_item em SolicitacaoRotaItem
        consulta_item = SolicitacaoRotaItem.objects.filter(ticket_item=ticket)
        if solicitacao_id:
            try:
                consulta_item = consulta_item.exclude(solicitacao_id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_item.exists():
            item_existente = consulta_item.first()
            if item_existente and item_existente.solicitacao:
                tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                tipo_existente = tipo_dict.get(item_existente.solicitacao.tipo, item_existente.solicitacao.tipo)
                
                return JsonResponse({
                    'existe': True,
                    'tipo': 'item_rota',
                    'message': f'O ID "{ticket}" já está cadastrado como item de rota em uma solicitação "{tipo_existente}".',
                    'detalhes': {
                        'tipo': tipo_existente,
                        'solicitacao_titulo': item_existente.solicitacao.titulo
                    }
                })
        
        # ID não encontrado
        return JsonResponse({
            'existe': False,
            'message': 'ID disponível'
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'existe': False,
            'message': f'Erro ao verificar ID: {str(e)}'
        }, status=500)

    
    try:
        # Buscar todas as solicitações recusadas não excluídas
        solicitacoes_recusadas = Solicitacoes.objects.filter(status='recusado', excluida=False)
        quantidade = solicitacoes_recusadas.count()
        
        if quantidade == 0:
            return JsonResponse({
                'success': True,
                'message': 'Não há solicitações recusadas para excluir.',
                'quantidade': 0
            })
        
        # Exclusão lógica em lote
        agora = timezone.now()
        for solicitacao in solicitacoes_recusadas:
            registrar_exclusao_solicitacao(
                solicitacao=solicitacao,
                usuario=request.user,
                origem=request.path,
                motivo="Exclusão em lote de solicitações recusadas",
            )
            solicitacao.excluida = True
            solicitacao.data_exclusao = agora
            solicitacao.excluida_por = request.user
            solicitacao.save(update_fields=['excluida', 'data_exclusao', 'excluida_por'])
        
        return JsonResponse({
            'success': True,
            'message': f'{quantidade} solicitação(ões) recusada(s) excluída(s) com sucesso!',
            'quantidade': quantidade
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao excluir solicitações: {str(e)}'
        }, status=500)


@require_http_methods(["GET"])
def verificar_id_existente(request):
    """
    View para verificar se um ID (ticket) já está cadastrado no sistema.
    Usado para validação em tempo real no formulário.
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'existe': False,
            'message': 'Não autenticado'
        }, status=401)
    
    ticket = request.GET.get('ticket', '').strip()
    solicitacao_id = request.GET.get('solicitacao_id', '').strip()  # Para modo edição
    
    if not ticket:
        return JsonResponse({
            'existe': False,
            'message': 'ID não fornecido'
        })
    
    try:
        # Verificar se existe como ticket principal em Solicitacoes
        consulta_solic = Solicitacoes.objects.filter(ticket=ticket)
        if solicitacao_id:
            try:
                consulta_solic = consulta_solic.exclude(id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_solic.exists():
            solicitacao_existente = consulta_solic.first()
            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
            tipo_existente = tipo_dict.get(solicitacao_existente.tipo, solicitacao_existente.tipo)
            status_existente = solicitacao_existente.status
            
            return JsonResponse({
                'existe': True,
                'tipo': 'solicitacao',
                'message': f'O ID "{ticket}" já está cadastrado como solicitação "{tipo_existente}" (Status: {status_existente}).',
                'detalhes': {
                    'tipo': tipo_existente,
                    'status': status_existente,
                    'titulo': solicitacao_existente.titulo
                }
            })
        
        # Verificar se existe como ticket_item em SolicitacaoRotaItem
        consulta_item = SolicitacaoRotaItem.objects.filter(ticket_item=ticket)
        if solicitacao_id:
            try:
                consulta_item = consulta_item.exclude(solicitacao_id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_item.exists():
            item_existente = consulta_item.first()
            if item_existente and item_existente.solicitacao:
                tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                tipo_existente = tipo_dict.get(item_existente.solicitacao.tipo, item_existente.solicitacao.tipo)
                
                return JsonResponse({
                    'existe': True,
                    'tipo': 'item_rota',
                    'message': f'O ID "{ticket}" já está cadastrado como item de rota em uma solicitação "{tipo_existente}".',
                    'detalhes': {
                        'tipo': tipo_existente,
                        'solicitacao_titulo': item_existente.solicitacao.titulo
                    }
                })
        
        # ID não encontrado
        return JsonResponse({
            'existe': False,
            'message': 'ID disponível'
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'existe': False,
            'message': f'Erro ao verificar ID: {str(e)}'
        }, status=500)

    
    try:
        # Buscar todas as solicitações recusadas não excluídas
        solicitacoes_recusadas = Solicitacoes.objects.filter(status='recusado', excluida=False)
        quantidade = solicitacoes_recusadas.count()
        
        if quantidade == 0:
            return JsonResponse({
                'success': True,
                'message': 'Não há solicitações recusadas para excluir.',
                'quantidade': 0
            })
        
        # Exclusão lógica em lote
        agora = timezone.now()
        for solicitacao in solicitacoes_recusadas:
            registrar_exclusao_solicitacao(
                solicitacao=solicitacao,
                usuario=request.user,
                origem=request.path,
                motivo="Exclusão em lote de solicitações recusadas",
            )
            solicitacao.excluida = True
            solicitacao.data_exclusao = agora
            solicitacao.excluida_por = request.user
            solicitacao.save(update_fields=['excluida', 'data_exclusao', 'excluida_por'])
        
        return JsonResponse({
            'success': True,
            'message': f'{quantidade} solicitação(ões) recusada(s) excluída(s) com sucesso!',
            'quantidade': quantidade
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'message': f'Erro ao excluir solicitações: {str(e)}'
        }, status=500)


@require_http_methods(["GET"])
def verificar_id_existente(request):
    """
    View para verificar se um ID (ticket) já está cadastrado no sistema.
    Usado para validação em tempo real no formulário.
    """
    if not request.user.is_authenticated:
        return JsonResponse({
            'existe': False,
            'message': 'Não autenticado'
        }, status=401)
    
    ticket = request.GET.get('ticket', '').strip()
    solicitacao_id = request.GET.get('solicitacao_id', '').strip()  # Para modo edição
    
    if not ticket:
        return JsonResponse({
            'existe': False,
            'message': 'ID não fornecido'
        })
    
    try:
        # Verificar se existe como ticket principal em Solicitacoes
        consulta_solic = Solicitacoes.objects.filter(ticket=ticket)
        if solicitacao_id:
            try:
                consulta_solic = consulta_solic.exclude(id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_solic.exists():
            solicitacao_existente = consulta_solic.first()
            tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
            tipo_existente = tipo_dict.get(solicitacao_existente.tipo, solicitacao_existente.tipo)
            status_existente = solicitacao_existente.status
            
            return JsonResponse({
                'existe': True,
                'tipo': 'solicitacao',
                'message': f'O ID "{ticket}" já está cadastrado como solicitação "{tipo_existente}" (Status: {status_existente}).',
                'detalhes': {
                    'tipo': tipo_existente,
                    'status': status_existente,
                    'titulo': solicitacao_existente.titulo
                }
            })
        
        # Verificar se existe como ticket_item em SolicitacaoRotaItem
        consulta_item = SolicitacaoRotaItem.objects.filter(ticket_item=ticket)
        if solicitacao_id:
            try:
                consulta_item = consulta_item.exclude(solicitacao_id=int(solicitacao_id))
            except (ValueError, TypeError):
                pass
        
        if consulta_item.exists():
            item_existente = consulta_item.first()
            if item_existente and item_existente.solicitacao:
                tipo_dict = dict(Solicitacoes.TIPO_CHOICES)
                tipo_existente = tipo_dict.get(item_existente.solicitacao.tipo, item_existente.solicitacao.tipo)
                
                return JsonResponse({
                    'existe': True,
                    'tipo': 'item_rota',
                    'message': f'O ID "{ticket}" já está cadastrado como item de rota em uma solicitação "{tipo_existente}".',
                    'detalhes': {
                        'tipo': tipo_existente,
                        'solicitacao_titulo': item_existente.solicitacao.titulo
                    }
                })
        
        # ID não encontrado
        return JsonResponse({
            'existe': False,
            'message': 'ID disponível'
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'existe': False,
            'message': f'Erro ao verificar ID: {str(e)}'
        }, status=500)
SolicitacaoTecnico
